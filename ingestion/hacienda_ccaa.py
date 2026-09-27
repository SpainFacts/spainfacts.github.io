"""Fuente dlt para las cuentas de las Comunidades Autónomas (Ministerio de Hacienda).

Liquidación de los presupuestos de las CCAA que publica la SGCIEF
(Subdirección General de Coordinación Financiera con las CCAA):
https://serviciostelematicosext.hacienda.gob.es/sgcief/publicacionliquidaciones/aspx/menuinicio.aspx

Se usa la sección "Descarga Datos Consolidados": cada fichero agrega la
Administración General de la comunidad con sus organismos autónomos, entes
y empresas incluidos en la consolidación, eliminando las transferencias
internas entre ellos (la lista está en "Organismos y Entidades incluidos
consolidación"). Es la cifra comparable entre comunidades; los datos "sin
consolidar" solo cubren la Administración General y no se cargan.

No hay API: son informes .xlsx pequeños, uno por comunidad y ejercicio, con
importes en MILES de euros. El portal es ASP.NET y exige una cookie de
sesión: primero se abre menuinicio.aspx con una requests.Session y después
las descargas funcionan (si no, redirige a sesion_expirada.htm).

Recursos (carga completa "replace", ~1.000 descargas y unos minutos):

1. `hacienda_ccaa_capitulos` (DescargaEconomicaDC.aspx): ingresos y gastos
   por capítulo económico (1-9) con presupuesto inicial, definitivo,
   derechos/obligaciones reconocidas netas y cobros/pagos líquidos.
2. `hacienda_ccaa_funcional` (DescargaFuncionalDC.aspx): obligaciones
   reconocidas por área (1 dígito) y política de gasto (2 dígitos), brutas y
   "depuradas IFL y PAC" (sin la participación de las entidades locales en
   tributos del Estado ni los fondos agrícolas europeos, que la comunidad solo
   hace pasar y duplicarían gasto de otras administraciones).
   Ojo: hasta ~2009 (según la comunidad) el fichero viene en la antigua
   clasificación por "grupos de función y funciones", con códigos que NO
   casan con las áreas/políticas actuales (41 era Sanidad, hoy es
   Agricultura). Se guarda la columna `clasificacion` para distinguirlas.

Códigos de comunidad (`cod_ccaa_hacienda`): 00 total, 01-19 en numeración de
Hacienda, que difiere de la del INE en 10-17 (Hacienda: 10 Extremadura,
11 Galicia, 12 Madrid, 13 Murcia, 14 Navarra, 15 País Vasco, 16 La Rioja,
17 C. Valenciana). La traducción al código INE se hace en dbt con la
semilla transform/seeds/ccaa_codigos_hacienda.csv. Ceuta (18) y Melilla (19)
están en la publicación como ciudades autónomas.
"""

import io
import logging
import re
import time
import unicodedata
from datetime import date

import dlt
import requests
from openpyxl import load_workbook

BASE = "https://serviciostelematicosext.hacienda.gob.es/sgcief/publicacionliquidaciones/aspx/"
CABECERAS = {"User-Agent": "SpainFacts/1.0 (+https://spainfacts.github.io)"}
PRIMER_ANIO = 2002
# 00 = total comunidades (sirve para cuadrar); 01-19 en numeración de Hacienda
COMUNIDADES = [f"{i:02d}" for i in range(0, 20)]
PAUSA_S = 0.3  # cortesía entre peticiones
TOLERANCIA_MILES = 1.0  # descuadre admitido entre suma de capítulos y total (redondeos)

log = logging.getLogger(__name__)

# Etiqueta con código: "1. Impuestos Directos", "31. Sanidad"
RE_CODIGO = re.compile(r"^\s*(\d{1,2})\.\s*(.+?)\s*$")


def _normaliza(texto) -> str:
    """Minúsculas, sin tildes y con espacios simples, para localizar cabeceras por texto."""
    t = unicodedata.normalize("NFKD", str(texto or "")).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", t).strip().lower()


# Cabecera de columna del fichero económico -> medida canónica
MEDIDAS_ECONOMICA = {
    "presupuesto inicial": "presupuesto_inicial",
    "presupuesto definitivo": "presupuesto_definitivo",
    "derechos reconocidos netos": "ejecutado",
    "obligaciones reconocidas netas": "ejecutado",
    "ingresos liquidos": "cobrado_pagado",
    "pagos liquidos": "cobrado_pagado",
}


def _medida_economica(cabecera) -> str | None:
    n = _normaliza(cabecera)
    if not n:
        return None
    return MEDIDAS_ECONOMICA.get(n, re.sub(r"[^a-z0-9]+", "_", n).strip("_"))


def _medida_funcional(cabecera) -> str | None:
    n = _normaliza(cabecera)
    if not n or "%" in n:  # los porcentajes se recalculan si hacen falta
        return None
    if "depurad" in n:
        return "obligaciones_depuradas"
    if "obligaciones" in n:
        return "obligaciones_brutas"
    return re.sub(r"[^a-z0-9]+", "_", n).strip("_")


def _num(valor) -> float | None:
    if valor in (None, ""):
        return None
    if isinstance(valor, (int, float)):
        return float(valor)
    return float(str(valor).replace(".", "").replace(",", "."))


class _Portal:
    """Sesión única con el portal de Hacienda: cookie ASP.NET, reintentos y pausa."""

    def __init__(self):
        self.s = None
        self._abre_sesion()

    def _abre_sesion(self):
        self.s = requests.Session()
        self.s.headers.update(CABECERAS)
        self.s.get(BASE + "menuinicio.aspx", timeout=60).raise_for_status()

    def ejercicios(self) -> list[int]:
        """Ejercicios publicados según el desplegable de descargas (el último cambia cada año)."""
        try:
            r = self.s.get(BASE + "SelDescargaDC.aspx", timeout=60)
            r.raise_for_status()
            bloque = re.search(r'name="ctl00\$MainContent\$ano".*?</select>', r.text, re.S)
            anios = sorted({int(a) for a in re.findall(r'value="(\d{4})"', bloque.group(0))})
            if anios:
                return [a for a in anios if a >= PRIMER_ANIO]
        except Exception as e:  # noqa: BLE001 - si falla, se prueba año a año
            log.warning("No se pudo leer la lista de ejercicios (%s); se usa un rango fijo", e)
        return list(range(PRIMER_ANIO, date.today().year))

    def xlsx(self, pagina: str, cod: str, anio: int):
        """Descarga un informe y devuelve la hoja activa (o None si no hay fichero)."""
        url = f"{BASE}{pagina}.aspx?cdcdad={cod}&ano={anio}"
        expiradas = 0
        for intento in range(6):
            time.sleep(PAUSA_S if intento == 0 else 2**intento)
            try:
                r = self.s.get(url, timeout=90)
            except requests.RequestException as e:
                log.warning("%s: %s (intento %d)", url, e, intento + 1)
                self._abre_sesion_segura()
                continue
            # El portal manda a sesion_expirada.htm tanto si caduca la cookie como
            # si el informe no existe (error de servidor, p. ej. Ceuta 2002 funcional):
            # se reabre la sesión una vez y, si se repite, se da por "sin fichero".
            if "sesion_expirada" in r.url:
                expiradas += 1
                if expiradas >= 2:
                    log.warning("%s: sin fichero (el portal devuelve error)", url)
                    return None
                self._abre_sesion_segura()
                continue
            if r.status_code >= 500 or r.status_code == 429:
                log.warning("%s: HTTP %s (intento %d)", url, r.status_code, intento + 1)
                continue
            r.raise_for_status()
            if r.content[:2] == b"PK":  # un .xlsx es un zip
                return load_workbook(io.BytesIO(r.content), data_only=True, read_only=True).active
            log.warning("%s no devolvió un .xlsx (%s)", url, r.headers.get("content-type"))
            return None
        raise RuntimeError(f"No se pudo descargar {url}")

    def _abre_sesion_segura(self):
        try:
            self._abre_sesion()
        except requests.RequestException as e:
            log.warning("No se pudo reabrir la sesión: %s", e)


def _filas(hoja) -> list[tuple]:
    return [fila for fila in hoja.iter_rows(values_only=True) if any(v not in (None, "") for v in fila)]


def _parsea_economica(hoja, cod: str, anio: int):
    """Capítulos de ingresos y de gastos. Las cabeceras se localizan por texto."""
    tipo, medidas, suma, filas = None, {}, {}, []
    for fila in _filas(hoja):
        etiqueta = _normaliza(fila[0])
        if etiqueta.startswith("capitulos de ingresos") or etiqueta.startswith("capitulos de gastos"):
            tipo = "I" if "ingresos" in etiqueta else "G"
            medidas = {j: _medida_economica(v) for j, v in enumerate(fila) if j > 0 and _medida_economica(v)}
            continue
        if tipo is None:
            continue
        m = RE_CODIGO.match(str(fila[0] or ""))
        if m:
            capitulo = int(m.group(1))
            for j, medida in medidas.items():
                importe = _num(fila[j]) if j < len(fila) else None
                if importe is None:
                    continue
                suma[(tipo, medida)] = suma.get((tipo, medida), 0.0) + importe
                filas.append({
                    "cod_ccaa_hacienda": cod, "anio": anio, "tipo": tipo,
                    "capitulo": capitulo, "capitulo_nombre": m.group(2),
                    "medida": medida, "importe_miles": importe,
                })
        elif etiqueta.startswith("total"):
            # Comprobación: la suma de capítulos debe cuadrar con el total del propio fichero
            for j, medida in medidas.items():
                total = _num(fila[j]) if j < len(fila) else None
                calculado = suma.get((tipo, medida), 0.0)
                if total is not None and abs(total - calculado) > TOLERANCIA_MILES:
                    log.warning("Descuadre capítulos %s %s %s %s: total %s vs suma %s",
                                cod, anio, tipo, medida, total, calculado)
            tipo = None
    return filas


def _parsea_funcional(hoja, cod: str, anio: int):
    """Áreas (1 dígito) y políticas (2 dígitos) de gasto, o grupos de función y funciones."""
    clasificacion, medidas, suma, filas = None, {}, {}, []
    for fila in _filas(hoja):
        etiqueta = _normaliza(fila[0])
        if etiqueta.startswith("areas y politicas"):
            clasificacion = "politicas"
        elif etiqueta.startswith("grupos de funcion"):
            clasificacion = "funciones"
        if clasificacion and not medidas and (etiqueta.startswith("areas") or etiqueta.startswith("grupos")):
            medidas = {j: _medida_funcional(v) for j, v in enumerate(fila) if j > 0 and _medida_funcional(v)}
            continue
        if not medidas:
            continue
        m = RE_CODIGO.match(str(fila[0] or ""))
        if m:
            codigo = m.group(1)
            nivel = "area" if len(codigo) == 1 else "politica"
            for j, medida in medidas.items():
                importe = _num(fila[j]) if j < len(fila) else None
                if importe is None:
                    continue
                suma[(nivel, medida)] = suma.get((nivel, medida), 0.0) + importe
                filas.append({
                    "cod_ccaa_hacienda": cod, "anio": anio, "clasificacion": clasificacion,
                    "nivel": nivel, "codigo": codigo, "nombre": m.group(2),
                    "cod_area": codigo[0], "medida": medida, "importe_miles": importe,
                })
        elif etiqueta.startswith("total"):
            for j, medida in medidas.items():
                total = _num(fila[j]) if j < len(fila) else None
                for nivel in ("area", "politica"):
                    calculado = suma.get((nivel, medida), 0.0)
                    if total is not None and abs(total - calculado) > TOLERANCIA_MILES:
                        log.warning("Descuadre funcional %s %s %s %s: total %s vs suma %s",
                                    cod, anio, nivel, medida, total, calculado)
            break
    return filas


@dlt.resource(name="hacienda_ccaa_capitulos", write_disposition="replace")
def capitulos(portal: _Portal, anios: list[int]):
    for anio in anios:
        for cod in COMUNIDADES:
            hoja = portal.xlsx("DescargaEconomicaDC", cod, anio)
            if hoja is not None:
                yield _parsea_economica(hoja, cod, anio)


@dlt.resource(name="hacienda_ccaa_funcional", write_disposition="replace")
def funcional(portal: _Portal, anios: list[int]):
    for anio in anios:
        for cod in COMUNIDADES:
            hoja = portal.xlsx("DescargaFuncionalDC", cod, anio)
            if hoja is not None:
                yield _parsea_funcional(hoja, cod, anio)


@dlt.source(name="hacienda_ccaa")
def hacienda_ccaa(anios: list[int] | None = None):
    portal = _Portal()
    anios = anios or portal.ejercicios()
    return [capitulos(portal, anios), funcional(portal, anios)]
