"""Fuente dlt de las gasolineras y los carburantes (tema `carburantes`).

1) Geoportal de gasolineras del Ministerio para la Transición Ecológica, API REST abierta:
   https://sedeaplicaciones.minetur.gob.es/ServiciosRESTCarburantes/PreciosCarburantes/
   - EstacionesTerrestres/                 foto actual (~11.500 estaciones, 12 MB, precios cada 30 min)
   - EstacionesTerrestresHist/dd-mm-aaaa   el mismo listado con los precios de ese día; desde el
                                           01-01-2007 hasta ayer (8-12 MB y 2-4 s por día).
   Trampas: JSON con BOM (utf-8-sig); números como texto con coma decimal y "" si no hay dato;
   IDMunicipio e IDCCAA son códigos propios del Ministerio (IDCCAA va en orden alfabético, no
   el del INE): el municipio se asigna por coordenadas con el localizador de ingestion/recarga
   y la provincia con IDProvincia (ese sí es el INE). En el histórico, rótulo, horario y
   dirección son los de HOY (en 2007 salen estaciones «MOEVE», marca de 2024): de él solo se
   usan precios. Los listados exigen la barra final (sin ella, 307 a una página HTML).
   El servidor solo negocia cifrados TLS antiguos: con la configuración por defecto de Python
   corta la conexión (WinError 10054); se usa una sesión con DEFAULT@SECLEVEL=1 solo para él.

2) CORES, consumo de productos petrolíferos por provincia, mensual desde 1997 (toneladas):
   https://www.cores.es/sites/default/files/archivos/estadisticas/consumos-pp-ccaa-provincias.xlsx
   Provincias por nombre, mes en texto, filas de total intercaladas.

3) Comisión Europea, Weekly Oil Bulletin, histórico desde 2005 (precio de venta al público con
   y sin impuestos de cada país de la UE, € por 1.000 l):
   https://energy.ec.europa.eu/data-and-analysis/weekly-oil-bulletin_en
   El fichero cambia de URL de vez en cuando: si falla, se busca el enlace en esa página.

Recursos:
  carburantes_estaciones          foto actual, una fila por estación (replace)
  carburantes_precios_dia         precio medio y mediano por provincia de gasóleo A, gasolina 95
                                  y GLP los días 1, 8, 15 y 22 de cada mes desde 2007 (merge,
                                  incremental: solo pide los días posteriores al último cargado).
                                  Cada día ya resumido se guarda además en data/carburantes_cache/
                                  (la primera carga son ~950 días, unos 40 minutos).
  carburantes_consumo_provincia   CORES (replace)
  carburantes_precios_ue          Weekly Oil Bulletin por país y semana (replace)
"""

import io
import json
import logging
import re
import ssl
import statistics
import time
import unicodedata
from datetime import date, timedelta
from pathlib import Path

import dlt
import requests
import requests.adapters

log = logging.getLogger(__name__)

REPO_ROOT = Path(__file__).resolve().parent.parent
CACHE = REPO_ROOT / "data" / "carburantes_cache"
API = "https://sedeaplicaciones.minetur.gob.es/ServiciosRESTCarburantes/PreciosCarburantes/"
CORES = "https://www.cores.es/sites/default/files/archivos/estadisticas/consumos-pp-ccaa-provincias.xlsx"
WOB_PAGINA = "https://energy.ec.europa.eu/data-and-analysis/weekly-oil-bulletin_en"
WOB = (
    "https://energy.ec.europa.eu/document/download/906e60ca-8b6a-44e7-8589-652854d2fd3f_en"
    "?filename=Weekly_Oil_Bulletin_Prices_History_maticni_4web.xlsx"
)
CABECERAS = {"User-Agent": "spainfacts.org (datos abiertos; contacto en github.com/SpainFacts)", "Accept": "application/json"}
PRIMER_DIA = date(2007, 1, 1)
DIAS_DEL_MES = (1, 8, 15, 22)

# Columna del API -> clave corta
PRODUCTOS = {
    "Precio Gasoleo A": "gasoleo_a",
    "Precio Gasolina 95 E5": "gasolina_95",
    "Precio Gases licuados del petróleo": "glp",
}
OTROS_PRECIOS = {
    "Precio Gasoleo Premium": "gasoleo_premium",
    "Precio Gasolina 98 E5": "gasolina_98",
    "Precio Gas Natural Comprimido": "gnc",
    "Precio Gas Natural Licuado": "gnl",
    "Precio Hidrogeno": "hidrogeno",
    "Precio Diésel Renovable": "diesel_renovable",
    "Precio Adblue": "adblue",
}


class _CifradosAntiguos(requests.adapters.HTTPAdapter):
    def init_poolmanager(self, *args, **kwargs):
        ctx = ssl.create_default_context()
        ctx.set_ciphers("DEFAULT@SECLEVEL=1")
        kwargs["ssl_context"] = ctx
        return super().init_poolmanager(*args, **kwargs)


_SESION = requests.Session()
_SESION.mount("https://sedeaplicaciones.minetur.gob.es/", _CifradosAntiguos())


def _get(url: str, intentos: int = 4, **kw) -> requests.Response:
    for i in range(intentos):
        try:
            r = _SESION.get(url, headers=CABECERAS, timeout=120, **kw)
            if r.status_code < 500:
                r.raise_for_status()
                return r
        except requests.RequestException as e:
            if i == intentos - 1:
                raise
            log.warning("%s: %s; reintento", url, e)
        time.sleep(5 * (i + 1))
    r.raise_for_status()
    return r


def _num(v) -> float | None:
    if v is None:
        return None
    v = str(v).strip().replace(",", ".")
    if not v:
        return None
    try:
        return float(v)
    except ValueError:
        return None


def _json(r: requests.Response) -> dict:
    return json.loads(r.content.decode("utf-8-sig"))


def _rotulo(texto) -> str:
    """Rótulo limpio en mayúsculas sin tildes; los números de estación («Nº 10.935») = sin rótulo."""
    t = unicodedata.normalize("NFKD", str(texto or "")).encode("ascii", "ignore").decode().upper()
    t = re.sub(r"\s+", " ", t).strip(" .-()")
    if not t or t in ("SIN ROTULO", "SIN MARCA", "BLANCA", "LIBRE") or re.fullmatch(r"(N[O]?|NUM)?\s*[\dº.\s]+", t):
        return "SIN ROTULO"
    return t


# ---------------------------------------------------------------- 1) foto actual


def estaciones():
    from ingestion.recarga import Localizador

    datos = _json(_get(API + "EstacionesTerrestres/"))
    localizar = Localizador()
    fecha = datos.get("Fecha")
    for e in datos["ListaEESSPrecio"]:
        lat, lon = _num(e.get("Latitud")), _num(e.get("Longitud (WGS84)"))
        # Hay estaciones con 0,0 y alguna con latitud y longitud cambiadas
        if lat is not None and lon is not None and not 27 < lat < 44.5 and 27 < lon < 44.5:
            lat, lon = lon, lat
        if not lat or not lon or not 27 < lat < 44.5:
            lat = lon = None
        cod_mun = None
        if lat is not None and lon is not None:
            cod_mun, _ = localizar(lon, lat)
        fila = {
            "ideess": e["IDEESS"],
            "rotulo": _rotulo(e.get("Rótulo")),
            "rotulo_original": (e.get("Rótulo") or "").strip(),
            "cod_prov": e.get("IDProvincia"),
            "cod_mun": cod_mun,
            "municipio_fuente": (e.get("Municipio") or "").strip(),
            "localidad": (e.get("Localidad") or "").strip(),
            "direccion": (e.get("Dirección") or "").strip(),
            "codigo_postal": e.get("C.P."),
            "horario": e.get("Horario"),
            "abierta_24h": "24H" in (e.get("Horario") or "").upper(),
            "latitud": lat,
            "longitud": lon,
            "tipo_venta": (e.get("Tipo Venta") or "").upper() or None,
            "remision": e.get("Remisión"),
            "fecha_precios": fecha,
        }
        for col, clave in {**PRODUCTOS, **OTROS_PRECIOS}.items():
            fila[f"precio_{clave}"] = _num(e.get(col))
        yield fila


# ---------------------------------------------------------------- 2) histórico de precios


def dias_a_cargar(desde: date = PRIMER_DIA, hasta: date | None = None):
    hasta = hasta or date.today() - timedelta(days=1)
    m = desde.replace(day=1)
    while m <= hasta:
        for d in DIAS_DEL_MES:
            dia = m.replace(day=d)
            if desde <= dia <= hasta:
                yield dia
        m = (m.replace(day=28) + timedelta(days=4)).replace(day=1)


def resumen_dia(lista: list[dict]) -> list[dict]:
    """Media, mediana y nº de estaciones con precio de cada producto por provincia y en España."""
    grupos: dict[str, dict[str, list[float]]] = {}
    for e in lista:
        if (e.get("Tipo Venta") or "P").upper() != "P":
            continue  # venta restringida (cooperativas): no es precio de venta al público
        for prov in (e.get("IDProvincia"), "00"):
            if not prov:
                continue
            g = grupos.setdefault(prov, {k: [] for k in PRODUCTOS.values()})
            for col, clave in PRODUCTOS.items():
                p = _num(e.get(col))
                if p and 0.2 < p < 5:  # fuera precios imposibles (errores de tecleo)
                    g[clave].append(p)
    filas = []
    for prov, g in grupos.items():
        for clave, precios in g.items():
            if precios:
                filas.append({
                    "cod_prov": prov,
                    "producto": clave,
                    "estaciones": len(precios),
                    "precio_medio": statistics.fmean(precios),
                    "precio_mediano": statistics.median(precios),
                })
    return filas


def precios_dia(dia: date) -> list[dict]:
    ruta = CACHE / f"{dia.isoformat()}.json"
    if ruta.exists():
        return json.loads(ruta.read_text(encoding="utf-8"))
    r = _get(API + f"EstacionesTerrestresHist/{dia:%d-%m-%Y}")
    lista = _json(r).get("ListaEESSPrecio") or []
    filas = resumen_dia(lista)
    if filas:  # un día vacío (caída del servicio) no se guarda, para reintentarlo
        CACHE.mkdir(parents=True, exist_ok=True)
        ruta.write_text(json.dumps(filas), encoding="utf-8")
    return filas


def precios_historicos(desde: date = PRIMER_DIA):
    for dia in dias_a_cargar(desde):
        try:
            filas = precios_dia(dia)
        except requests.RequestException as e:
            log.warning("Sin precios de %s: %s", dia, e)
            continue
        for f in filas:
            yield {"fecha": dia, **f}


# ---------------------------------------------------------------- 3) CORES


MESES = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto",
         "septiembre", "octubre", "noviembre", "diciembre"]
COLUMNAS_CORES = {
    "gasolina 97": "gasolina_97_t",
    "gasolina 95": "gasolina_95_t",
    "gasolina 98": "gasolina_98_t",
    "gasoleo a": "gasoleo_a_t",
    "gasoleo b": "gasoleo_b_t",
    "gasoleo c": "gasoleo_c_t",
    "fueloleo": "fueloleo_t",
}


def _sin_tildes(t) -> str:
    return unicodedata.normalize("NFKD", str(t or "")).encode("ascii", "ignore").decode().lower().strip()


def consumo_provincia(contenido: bytes | None = None):
    import openpyxl

    contenido = contenido or _get(CORES).content
    hoja = openpyxl.load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)["Consumos"]
    cab = None
    for fila in hoja.iter_rows(values_only=True):
        if cab is None:
            if fila and _sin_tildes(fila[0]) == "ano":
                cab = {}
                for i, c in enumerate(fila):
                    n = _sin_tildes(c)
                    for pref, clave in COLUMNAS_CORES.items():
                        if n.startswith(pref):
                            cab[clave] = i
            continue
        anio, mes, ccaa, provincia = fila[0], _sin_tildes(fila[1]), fila[2], fila[3]
        if not isinstance(anio, (int, float)) or mes not in MESES:
            continue
        registro = {
            "anio": int(anio),
            "mes": MESES.index(mes) + 1,
            "ccaa": (str(ccaa).strip() if ccaa else None),
            "provincia": (str(provincia).strip() if provincia else None),
        }
        for clave, i in cab.items():
            v = fila[i] if i < len(fila) else None
            registro[clave] = float(v) if isinstance(v, (int, float)) else None
        yield registro


# ---------------------------------------------------------------- 4) Weekly Oil Bulletin


def _url_wob() -> str:
    try:
        _get(WOB, stream=True).close()
        return WOB
    except requests.RequestException:
        html = _get(WOB_PAGINA).text
        m = re.search(r'href="([^"]+Prices_History[^"]+\.xlsx[^"]*)"', html)
        if not m:
            raise
        url = m.group(1).replace("&amp;", "&")
        return url if url.startswith("http") else "https://energy.ec.europa.eu" + url


def precios_ue(contenido: bytes | None = None):
    """Una fila por semana, país, producto y con/sin impuestos (€ por litro)."""
    import openpyxl

    contenido = contenido or _get(_url_wob()).content
    libro = openpyxl.load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    patron = re.compile(r"^([A-Z]{2,3})_price_(with|wo)_tax_(euro95|diesel)$")
    for nombre_hoja in ("Prices with taxes", "Prices wo taxes"):
        hoja = libro[nombre_hoja]
        columnas = None
        for fila in hoja.iter_rows(values_only=True):
            if columnas is None:
                columnas = {i: patron.match(str(c or "").strip()) for i, c in enumerate(fila)}
                columnas = {i: m.groups() for i, m in columnas.items() if m}
                continue
            if not hasattr(fila[0], "year"):
                continue
            semana = fila[0].date()
            for i, (pais, imp, prod) in columnas.items():
                v = fila[i] if i < len(fila) else None
                if isinstance(v, (int, float)) and v > 0:
                    yield {
                        "semana": semana,
                        "cod_pais_fuente": pais,
                        "producto": "gasolina_95" if prod == "euro95" else "gasoleo_a",
                        "con_impuestos": imp == "with",
                        "precio_litro": v / 1000,
                    }


# ---------------------------------------------------------------- fuente


@dlt.source(name="carburantes", max_table_nesting=0)
def carburantes():
    @dlt.resource(name="carburantes_estaciones", write_disposition="replace")
    def r_estaciones():
        yield from estaciones()

    @dlt.resource(name="carburantes_precios_dia", write_disposition="merge", primary_key=("fecha", "cod_prov", "producto"))
    def r_precios(fecha=dlt.sources.incremental("fecha", initial_value=PRIMER_DIA)):
        # Solo los días posteriores al último cargado (el estado lo guarda dlt)
        ultimo = fecha.last_value or PRIMER_DIA
        yield from precios_historicos(ultimo + timedelta(days=1) if ultimo > PRIMER_DIA else PRIMER_DIA)

    @dlt.resource(name="carburantes_consumo_provincia", write_disposition="replace")
    def r_consumo():
        yield from consumo_provincia()

    @dlt.resource(name="carburantes_precios_ue", write_disposition="replace")
    def r_ue():
        yield from precios_ue()

    return r_estaciones, r_precios, r_consumo, r_ue
