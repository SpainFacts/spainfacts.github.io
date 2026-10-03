"""Fuente dlt para el sector de la construcción como actividad económica (página /economia/construccion).

Eurostat, API de diseminación (JSON-stat 2.0), los 27 países y el agregado EU27_2020:
  eurostat_construccion_vab        nama_10_a10: VAB (B1G) del total y de la construcción (rama F) en
                                   % del total (PC_TOT), precios corrientes (CP_MEUR) y volumen
                                   (CLV10_MEUR), desde 1995.
  eurostat_construccion_empleo     nama_10_a10_e: ocupados (EMP_DC, miles) del total y de la rama F.
  eurostat_construccion_poblacion  nama_10_pe: población media anual (miles), para los valores por habitante.
  eurostat_construccion_produccion sts_copr_a: índice de producción en la construcción (2021=100, corregido
                                   de calendario) de F, F41 (edificios), F42 (ingeniería civil) y F43
                                   (actividades especializadas), desde 1995. OJO: España no publica en I21 los
                                   agregados F_CC1/F_CC2 (edificación / obra civil por tipo de obra); se usan
                                   las ramas F41 y F42.
  eurostat_construccion_permisos   sts_cobp_a: permisos de construcción en viviendas (BPRM_DW, miles) y en
                                   superficie útil (BPRM_SQM, millones de m²): edificios residenciales
                                   (CPA_F41001), no residenciales (CPA_F41002) y total (CPA_F41001_41002); las
                                   viviendas solo vienen en CPA_F41001_X_410014 (residencial sin residencias
                                   colectivas).
  eurostat_construccion_costes     sts_copi_a: índice de costes de construcción (COST) y de precios de
                                   producción (PRC_PRR) de la vivienda nueva (CPA_F41001_X_410014), 2021=100.
  eurostat_construccion_sbs        sbs_ovw_act: empresas (ENT_NR), personas ocupadas (EMP_NR), cifra de
                                   negocios neta (NETTUR_MEUR), valor añadido (AV_MEUR) y productividad
                                   aparente (LABPRY_TEUR) de F, F41, F42 y F43, 2021-último.
  eurostat_construccion_vab_regional nama_10r_3gva: VAB total y de la rama F de las comunidades (NUTS 2) y de
                                   España, precios corrientes, desde 2000.
  eurostat_construccion_hicp       prc_hicp_aind: IPCA anual de España desde 1996, solo para enlazar el
                                   deflactor del INE (main.deflactor, desde 2002) hacia atrás.
INE (Tempus3), EPA (CNAE-2009; las tablas 79xxx de 2026 son la nueva CNAE-2025 y dan otras cifras):
  ine_construccion_ocupados_prov   65354: ocupados por sector económico y provincia, trimestral desde 2008.
  ine_construccion_parados_ccaa    65331: parados por sector económico (del último empleo) y comunidad,
                                   trimestral desde 2008.
Banco de España, Boletín Estadístico, capítulo 23 (be23.zip, CSV):
  construccion_bde_series          be2308 (visados de dirección de obra de los colegios de aparejadores,
                                   Ministerio de Transportes), be2309 (licitación oficial, Ministerio de
                                   Transportes), be2311 (cemento: producción, comercio exterior y consumo
                                   aparente, Oficemen / Ministerio de Industria), mensual.
ISTAC (Gobierno de Canarias), que redistribuye en abierto estadísticas del Ministerio de Transportes con
España y todas las comunidades (la web del ministerio bloquea las descargas automáticas):
  construccion_licitacion_ccaa     E20004A_000001: presupuesto de licitación oficial (miles de euros) por
                                   agente contratante (Estado y Seguridad Social / entes territoriales /
                                   AAPP) y tipo de obra (edificación / ingeniería civil), mensual y anual
                                   desde 1989.
  construccion_visados_ccaa        E20006A_000002: viviendas visadas (obra nueva, ampliación, reforma),
                                   España y comunidades, mensual y anual desde 2000 (sin Ceuta ni Melilla).
Seguridad Social, «Afiliación AAAA.xlsx» (hoja Tabla_4_ABC, antes DATOS PEST567: afiliados medios por
sección CNAE y provincia):
  construccion_afiliados           afiliados medios del mes de la sección F y del total, por provincia y
                                   régimen (General y Autónomos), desde 2021 (2026 ya en CNAE-2025).
"""

import csv
import io
import logging
import re

import dlt
import openpyxl
import requests

from ingestion.bde import _cuadros
from ingestion.eurostat import parse_json_stat_series
from ingestion.ine import _tabla_resource
from ingestion.pensiones import AFILIACION, _enlaces_xlsx, _get

log = logging.getLogger(__name__)

BASE = "https://ec.europa.eu/eurostat/api/dissemination/statistics/1.0/data/"
ISTAC = "https://datos.canarias.es/api/estadisticas/statistical-resources/v1.0/datasets/ISTAC/{id}/~latest.csv"

UE27 = ["AT", "BE", "BG", "CY", "CZ", "DE", "DK", "EE", "EL", "ES", "FI", "FR", "HR", "HU", "IE", "IT",
        "LT", "LU", "LV", "MT", "NL", "PL", "PT", "RO", "SE", "SI", "SK"]
NUTS2_ES = ["ES11", "ES12", "ES13", "ES21", "ES22", "ES23", "ES24", "ES30", "ES41", "ES42", "ES43",
            "ES51", "ES52", "ES53", "ES61", "ES62", "ES63", "ES64", "ES70"]

INE_TABLAS = {
    "ine_construccion_ocupados_prov": "65354",
    "ine_construccion_parados_ccaa": "65331",
}


def _json(url: str, timeout: int = 180) -> dict:
    resp = requests.get(url, timeout=timeout)
    resp.raise_for_status()
    return resp.json()


def _f(valor):
    # Eurostat mezcla enteros y decimales: float() para que dlt no cree columnas variantes
    return float(valor) if isinstance(valor, (int, float)) else None


def _geo(lista, campo="geo"):
    return "".join(f"&{campo}={g}" for g in lista)


def _paises():
    return _geo(UE27 + ["EU27_2020"])


def _istac(dataset: str, medida: str):
    url = ISTAC.format(id=dataset) + f"?dim=MEDIDAS:{medida}"
    resp = requests.get(url, timeout=300)
    resp.raise_for_status()
    texto = resp.content.decode("utf-8-sig")
    return csv.DictReader(io.StringIO(texto))


def _num(texto):
    texto = (texto or "").strip()
    if not texto:
        return None
    try:
        return float(texto)
    except ValueError:
        return None


def _periodo(codigo: str):
    # '2024' -> (2024, None); '2024-M07' -> (2024, 7)
    m = re.match(r"^(\d{4})(?:-M(\d{2}))?$", codigo.strip())
    if not m:
        return None, None
    return int(m.group(1)), (int(m.group(2)) if m.group(2) else None)


@dlt.source(name="construccion")
def construccion():
    @dlt.resource(name="eurostat_construccion_vab", write_disposition="replace")
    def vab():
        url = (BASE + "nama_10_a10?na_item=B1G&unit=PC_TOT&unit=CP_MEUR&unit=CLV10_MEUR"
               "&nace_r2=TOTAL&nace_r2=F&sinceTimePeriod=1995" + _paises() + "&format=JSON&lang=EN")
        for c, v in parse_json_stat_series(_json(url)):
            yield {"anio": int(c["time"]), "pais": c["geo"], "rama": c["nace_r2"], "unidad": c["unit"], "valor": _f(v)}

    @dlt.resource(name="eurostat_construccion_empleo", write_disposition="replace")
    def empleo():
        url = (BASE + "nama_10_a10_e?na_item=EMP_DC&unit=THS_PER&nace_r2=TOTAL&nace_r2=F"
               "&sinceTimePeriod=1995" + _paises() + "&format=JSON&lang=EN")
        for c, v in parse_json_stat_series(_json(url)):
            yield {"anio": int(c["time"]), "pais": c["geo"], "rama": c["nace_r2"], "miles": _f(v)}

    @dlt.resource(name="eurostat_construccion_poblacion", write_disposition="replace")
    def poblacion():
        url = (BASE + "nama_10_pe?na_item=POP_NC&unit=THS_PER&sinceTimePeriod=1995" + _paises()
               + "&format=JSON&lang=EN")
        for c, v in parse_json_stat_series(_json(url)):
            yield {"anio": int(c["time"]), "pais": c["geo"], "miles": _f(v)}

    @dlt.resource(name="eurostat_construccion_produccion", write_disposition="replace")
    def produccion():
        url = (BASE + "sts_copr_a?indic_bt=PRD&unit=I21&s_adj=CA&sinceTimePeriod=1995" + _paises()
               + _geo(["F", "F41", "F42", "F43"], "nace_r2") + "&format=JSON&lang=EN")
        for c, v in parse_json_stat_series(_json(url)):
            yield {"anio": int(c["time"]), "pais": c["geo"], "rama": c["nace_r2"], "indice": _f(v)}

    @dlt.resource(name="eurostat_construccion_permisos", write_disposition="replace")
    def permisos():
        url = (BASE + "sts_cobp_a?s_adj=NSA&unit=THS&unit=MIO_M2&sinceTimePeriod=1995" + _paises()
               + _geo(["CPA_F41001_41002", "CPA_F41001", "CPA_F41002", "CPA_F41001_X_410014"], "cpa2_1")
               + "&format=JSON&lang=EN")
        for c, v in parse_json_stat_series(_json(url)):
            yield {"anio": int(c["time"]), "pais": c["geo"], "edificio": c["cpa2_1"],
                   "indicador": c["indic_bt"], "unidad": c["unit"], "valor": _f(v)}

    @dlt.resource(name="eurostat_construccion_costes", write_disposition="replace")
    def costes():
        url = (BASE + "sts_copi_a?unit=I21&s_adj=NSA&cpa2_1=CPA_F41001_X_410014&sinceTimePeriod=1995"
               + _paises() + "&format=JSON&lang=EN")
        for c, v in parse_json_stat_series(_json(url)):
            yield {"anio": int(c["time"]), "pais": c["geo"], "indicador": c["indic_bt"], "indice": _f(v)}

    @dlt.resource(name="eurostat_construccion_sbs", write_disposition="replace")
    def sbs():
        indicadores = ["ENT_NR", "EMP_NR", "NETTUR_MEUR", "AV_MEUR", "LABPRY_TEUR"]
        url = (BASE + "sbs_ovw_act?" + _paises()[1:] + _geo(["F", "F41", "F42", "F43"], "nace_r2")
               + _geo(indicadores, "indic_sbs") + "&format=JSON&lang=EN")
        for c, v in parse_json_stat_series(_json(url)):
            yield {"anio": int(c["time"]), "pais": c["geo"], "rama": c["nace_r2"],
                   "indicador": c["indic_sbs"], "valor": _f(v)}

    @dlt.resource(name="eurostat_construccion_vab_regional", write_disposition="replace")
    def vab_regional():
        url = (BASE + "nama_10r_3gva?unit=CP_MEUR&nace_r2=TOTAL&nace_r2=F&sinceTimePeriod=2000"
               + _geo(NUTS2_ES + ["ES"]) + "&format=JSON&lang=EN")
        for c, v in parse_json_stat_series(_json(url)):
            yield {"anio": int(c["time"]), "nuts": c["geo"], "rama": c["nace_r2"], "mill_eur": _f(v)}

    @dlt.resource(name="eurostat_construccion_hicp", write_disposition="replace")
    def hicp():
        """prc_hicp_aind: IPCA anual de España (2015=100), para alargar el deflactor antes de 2002."""
        url = BASE + "prc_hicp_aind?geo=ES&coicop=CP00&unit=INX_A_AVG&format=JSON&lang=EN"
        for c, v in parse_json_stat_series(_json(url)):
            yield {"anio": int(c["time"]), "pais": c["geo"], "indice": _f(v)}

    @dlt.resource(name="construccion_bde_series", write_disposition="replace")
    def bde_series():
        for f in _cuadros("be23", ("be2308", "be2309", "be2311")):
            yield {
                "cuadro": f["cuadro"],
                "serie": f["serie"],
                "descripcion": (f["descripcion"] or "").replace("Estadísticas Generales. ", "")
                .replace("Estadísticas generales. ", ""),
                "unidad": f["unidad"],
                "anio": f["fecha"].year,
                "mes": f["fecha"].month,
                "valor": float(f["valor"]),
            }

    @dlt.resource(name="construccion_licitacion_ccaa", write_disposition="replace")
    def licitacion_ccaa():
        for r in _istac("E20004A_000001", "PRESUPUESTO_LICITACION"):
            anio, mes = _periodo(r["TIME_PERIOD_CODE"])
            if anio is None:
                continue
            yield {
                "nuts": r["TERRITORIO_CODE"],
                "territorio": r["TERRITORIO#es"],
                "anio": anio,
                "mes": mes,
                "agente": r["AGENTES_CONTRATANTES_CODE"],
                "tipo_obra": r["TIPOLOGIAS_OBRA_CODE"],
                "miles_eur": _num(r["OBS_VALUE"]),
                "estado": r.get("ESTADO_OBSERVACION#es") or None,
            }

    @dlt.resource(name="construccion_visados_ccaa", write_disposition="replace")
    def visados_ccaa():
        for r in _istac("E20006A_000002", "VIVIENDAS"):
            anio, mes = _periodo(r["TIME_PERIOD_CODE"])
            if anio is None:
                continue
            yield {
                "territorio": r["TERRITORIO#es"],
                "territorio_code": r["TERRITORIO_CODE"],
                "anio": anio,
                "mes": mes,
                "tipo_obra": r["TIPOS_OBRA_VIVIENDAS_CODE"],
                "tipo_obra_nombre": r["TIPOS_OBRA_VIVIENDAS#es"],
                "viviendas": _num(r["OBS_VALUE"]),
            }

    @dlt.resource(name="construccion_afiliados", write_disposition="replace")
    def afiliados():
        for nombre, enlace in _enlaces_xlsx(_get(AFILIACION).text):
            if not re.search(r"afiliaci.n\s*20\d\d", nombre, re.I):
                continue
            wb = openpyxl.load_workbook(io.BytesIO(_get(enlace).content), read_only=True, data_only=True)
            # la hoja de actividades por provincia se llama Tabla_4_ABC desde 2025 y DATOS PEST567 antes
            hoja = next((h for h in ("Tabla_4_ABC", "DATOS PEST567") if h in wb.sheetnames), None)
            if hoja is None:
                log.warning("construccion: %s sin hoja de actividades por provincia, se omite", nombre)
                continue
            filas = wb[hoja].iter_rows(values_only=True)
            cab = [str(x or "").upper() for x in next(filas)]
            ip, ir, ipr, isec, isal = (cab.index("PERIODO"), cab.index("DESREGIMEN"), cab.index("PROVINCIA"),
                                       cab.index("CD_SECCION"), cab.index("SALDOS"))
            suma = {}
            for f in filas:
                if f[ip] is None or not isinstance(f[isal], (int, float)) or f[ipr] is None:
                    continue
                periodo = str(f[ip]).strip()[:6]
                regimen = "Autónomos" if "AUTONOM" in str(f[ir]).upper() else "General"
                base = (int(periodo[:4]), int(periodo[4:6]), f"{int(f[ipr]):02d}", regimen)
                suma[base + ("Total",)] = suma.get(base + ("Total",), 0.0) + float(f[isal])
                if str(f[isec]).strip().upper() == "F":
                    suma[base + ("F",)] = suma.get(base + ("F",), 0.0) + float(f[isal])
            log.info("construccion: %s -> %d filas de afiliados", nombre, len(suma))
            cnae = "CNAE-2025" if "CNAE25" in nombre.upper() else "CNAE-2009"
            for (anio, mes, cod_prov, regimen, seccion), valor in suma.items():
                yield {"anio": anio, "mes": mes, "cod_prov": cod_prov, "regimen": regimen, "seccion": seccion,
                       "afiliados": valor, "cnae": cnae, "fichero": nombre}

    recursos = [vab, empleo, poblacion, produccion, permisos, costes, sbs, vab_regional, hicp,
                bde_series, licitacion_ccaa, visados_ccaa, afiliados]
    recursos += [_tabla_resource(nombre, tabla) for nombre, tabla in INE_TABLAS.items()]
    return recursos
