"""Fuente dlt del tema «sanidad pública y privada»: conciertos, seguros privados y gasto de los hogares.

Ministerio de Sanidad
  sanidad_privada_egsp        Estadística de Gasto Sanitario Público (EGSP, cuentas satélite, XLS
                              egspGastoReal, el mismo que lee ingestion/sanidad.py): gasto consolidado
                              de cada comunidad autónoma (tablas 3.4.01 a 3.4.17, código INE 01-17) y
                              del sector Comunidades Autónomas (tabla 2.4, cod_ccaa '00') en miles de
                              euros, 2002-último año (los dos últimos, provisionales), en cuatro
                              bloques: 'cuenta_satelite' (conceptos de la cuenta satélite: producción
                              pública por función y factor, producción de mercado comprada por la
                              administración —servicios hospitalarios y especializados concertados,
                              farmacia...—, transferencias y capital), 'economica' (clasificación
                              económico-presupuestaria: remuneración del personal, consumo intermedio,
                              consumo de capital fijo, conciertos, transferencias corrientes, gasto de
                              capital), 'funcional' (servicios hospitalarios y especializados,
                              primaria, salud pública, servicios colectivos, farmacia, traslado y
                              prótesis, capital) y 'contabilidad_nacional'.
  sanidad_privada_cnh         Catálogo Nacional de Hospitales (un XLSX por edición desde 2020; la
                              edición N describe la situación a 31-12-N-1): cada hospital con
                              comunidad, municipio, camas instaladas, clase, dependencia funcional
                              (servicio de salud, otros públicos, privado, mutua, ONG...) y concierto
                              con el SNS (sin concierto, parcial, sustitutorio, red de utilización
                              pública); 'DC' (dato del complejo) se sustituye por el del complejo.

INE (ficheros PC-Axis de INEbase en CSV, no están en la API Tempus)
  ine_sp_cobertura            Encuesta Nacional de Salud (ENSE 2006, 2011-12, 2017; cifras relativas)
                              y Encuesta Europea de Salud en España (EESE 2020; miles de personas,
                              población de 15 y más años): «modalidad de la cobertura sanitaria
                              (exclusiva)» por sexo y comunidad: pública exclusivamente, privada
                              exclusivamente, mixta (pública y seguro privado), otras situaciones.
INE (API Tempus)
  ine_sp_epf                  Encuesta de Presupuestos Familiares, grupo de gasto 06 «Sanidad» por
                              comunidad: gasto medio por hogar, por persona y % del gasto del hogar.
                              Dos series: tabla 73991 (clasificación ECOICOP actual, 2016-último año)
                              y tabla 28486 (base 2006, 2006-2023). Solo se guardan las series de
                              Sanidad (dato base).
"""

import csv
import io
import logging
import re

import dlt
import requests

from ingestion.sanidad import EGSP_XLS, UA, _cod_ccaa

log = logging.getLogger(__name__)

INE_TEMPUS = "https://servicios.ine.es/wstempus/js/ES/DATOS_TABLA/"
INE_PX = "https://www.ine.es/jaxi/files/_px/es/csv_bdsc/t15/p419/"
INE_TPX = "https://www.ine.es/jaxi/files/tpx/es/csv_bdsc/"

# (encuesta, año de referencia, URL del CSV, unidad, ámbito de población)
COBERTURA = [
    ("ENSE 2006", 2006, INE_PX + "a2006/p08/l0/04163.csv_bdsc", "pct", "Todas las edades"),
    ("ENSE 2011-2012", 2011, INE_PX + "a2011/p05/l0/05183.csv_bdsc", "pct", "Todas las edades"),
    ("ENSE 2017", 2017, INE_PX + "a2017/p05/l0/06004.csv_bdsc", "pct", "Todas las edades"),
    ("EESE 2020", 2020, INE_TPX + "47450.csv", "miles", "15 y más años"),
]

EPF_TABLAS = {"73991": "ecoicop", "28486": "base_2006"}

# Catálogo Nacional de Hospitales: un XLSX por edición desde 2020 (las anteriores son bases Access).
# La edición N recoge la situación a 31 de diciembre de N-1.
CNH_URL = ("https://www.sanidad.gob.es/estadEstudios/estadisticas/sisInfSanSNS/ofertaRecursos/"
           "hospitales/docs/CNH_{}.xlsx")
CNH_PRIMERA = 2020
CNH_ULTIMA = 2030  # las ediciones aún no publicadas dan 404 y se saltan

# Años de cabecera: '2015', '2023(*)' (provisional) o con letra de nota al pie ('2009a', '2020b').
ANIO = re.compile(r"^(\d{4})\s*[a-z]?\s*(\(\*\))?\s*[a-z]?$")
BLOQUES = [
    ("CUENTAS SAT", "cuenta_satelite"),
    ("CLASIFICACIÓN  ECONÓMICO", "economica"),
    ("CLASIFICACIÓN ECONÓMICO", "economica"),
    ("CLASIFICACIÓN  FUNCIONAL", "funcional"),
    ("CLASIFICACIÓN FUNCIONAL", "funcional"),
    ("AGREGADOS DE CONTABILIDAD NACIONAL", "contabilidad_nacional"),
    ("GASTO CONTABLE Y APORTACI", None),  # agentes de gasto: no se leen
]
SUBCONCEPTO = re.compile(r"^(\d+(\.\d+)*)\s*-\s*")


def _texto(v):
    if isinstance(v, float):
        return str(int(v)) if v == int(v) else str(v)
    return re.sub(r"\s+", " ", str(v)).strip()


def _hoja_egsp(sh, cod_ccaa):
    bloque, anios, padre, orden = None, {}, None, 0
    for f in range(sh.nrows):
        celdas = [_texto(sh.cell_value(f, c)) for c in range(sh.ncols)]
        cabecera = " ".join(celdas[:3]).upper()
        nuevo = next((b for clave, b in BLOQUES if clave in cabecera), "sin_cambio")
        if nuevo != "sin_cambio" and not any(isinstance(sh.cell_value(f, c), float) for c in range(3, sh.ncols)):
            bloque, anios, padre, orden = nuevo, {}, None, 0
            continue
        if bloque is None:
            continue
        cols_anio = {}
        for c in range(3, sh.ncols):
            m = ANIO.match(celdas[c])
            if m:
                cols_anio[c] = (int(m.group(1)), bool(m.group(2)))
        # Fila de años: al menos 10 años seguidos (1990-2050). Una fila de datos con importes de
        # cuatro cifras (p. ej. 2345,0) no debe confundirse con una cabecera.
        seguidos = [a for a, _ in cols_anio.values()]
        if len(seguidos) >= 10 and all(1990 <= a <= 2050 for a in seguidos) \
                and seguidos == list(range(seguidos[0], seguidos[0] + len(seguidos))):
            anios = cols_anio
            continue
        etiqueta = next((x for x in celdas[:3] if x), "")
        if not etiqueta or not anios or etiqueta.startswith("("):
            continue
        valores = {c: sh.cell_value(f, c) for c in anios}
        if not any(isinstance(v, float) for v in valores.values()):
            continue
        orden += 1
        es_sub = bloque == "cuenta_satelite" and etiqueta in (
            "Remuneración del personal", "Consumo intermedio", "Consumo de capital fijo")
        if bloque == "cuenta_satelite" and not es_sub:
            padre = etiqueta if SUBCONCEPTO.match(etiqueta) else None
        for c, (anio, provisional) in anios.items():
            v = valores[c]
            if isinstance(v, float):
                yield {"cod_ccaa": cod_ccaa, "hoja": sh.name, "clasificacion": bloque, "orden": orden,
                       "concepto": etiqueta, "padre": padre if es_sub else None,
                       "anio": anio, "provisional": provisional, "miles_eur": float(v)}


def _egsp():
    import xlrd

    r = requests.get(EGSP_XLS, headers=UA, timeout=300)
    r.raise_for_status()
    wb = xlrd.open_workbook(file_contents=r.content)
    hojas = [("Tabla 2.4", "00")] + [(f"Tabla 3.4.{i:02d}", f"{i:02d}") for i in range(1, 18)]
    for nombre, cod in hojas:
        n = 0
        for fila in _hoja_egsp(wb.sheet_by_name(nombre), cod):
            n += 1
            yield fila
        log.info("EGSP %s (%s): %d filas", nombre, cod, n)
        if n == 0:
            raise ValueError(f"EGSP: la hoja {nombre} no tiene datos (¿ha cambiado el formato?)")


def _numero(txt):
    txt = txt.strip()
    if not txt or txt in ("..", "-", "."):
        return None
    return float(txt.replace(".", "").replace(",", "."))


def _cobertura():
    for encuesta, anio, url, unidad, ambito in COBERTURA:
        r = requests.get(url, headers=UA, timeout=180)
        r.raise_for_status()
        texto = r.content.decode("utf-8-sig")
        lector = csv.reader(io.StringIO(texto), delimiter=";")
        next(lector)
        n = 0
        for fila in lector:
            if len(fila) < 5:
                continue
            sexo, _, territorio, cobertura, valor = fila[:5]
            territorio = re.sub(r"^\d{2}\s+", "", territorio.strip())
            cod = "00" if not territorio else _cod_ccaa(territorio)
            if cod is None:
                log.warning("Cobertura %s: territorio sin código %r", encuesta, territorio)
                continue
            n += 1
            yield {"encuesta": encuesta, "anio": anio, "ambito_poblacion": ambito, "unidad": unidad,
                   "sexo": sexo.strip(), "territorio": territorio or "Total Nacional", "cod_ccaa": cod,
                   "cobertura": cobertura.strip(), "valor": _numero(valor), "url": url}
        log.info("Cobertura %s: %d filas", encuesta, n)


def _epf():
    for tabla, serie_epf in EPF_TABLAS.items():
        r = requests.get(INE_TEMPUS + tabla, params={"nult": 40}, timeout=300)
        r.raise_for_status()
        for serie in r.json():
            nombre = serie.get("Nombre") or ""
            partes = [p.strip() for p in nombre.split(".") if p.strip()]
            if "Sanidad" not in partes or "Dato base" not in partes:
                continue
            medida = next((p for p in partes if p.startswith("Gasto medio") or p.startswith("Distribución")
                           or p == "Gasto total"), None)
            territorio = next((p for p in partes if p == "Total Nacional" or _cod_ccaa(p)), None)
            for punto in serie.get("Data", []):
                yield {"tabla": tabla, "serie_epf": serie_epf, "cod_serie": serie.get("COD"),
                       "serie": nombre, "territorio": territorio,
                       "cod_ccaa": "00" if territorio == "Total Nacional" else _cod_ccaa(territorio or ""),
                       "medida": medida, "anio": punto.get("Anyo"),
                       "valor": None if punto.get("Valor") is None else float(punto["Valor"]),
                       "secreto": punto.get("Secreto")}


def _cnh_hoja(wb, prefijo):
    nombre = next(n for n in wb.sheetnames if n.strip().upper().startswith(prefijo))
    filas = wb[nombre].iter_rows(values_only=True)
    cabecera = [re.sub(r"\s+", " ", str(c or "")).strip() for c in next(filas)]
    for fila in filas:
        if fila and any(fila):
            yield dict(zip(cabecera, fila))


def _cnh():
    import openpyxl

    for edicion in range(CNH_PRIMERA, CNH_ULTIMA + 1):
        url = CNH_URL.format(edicion)
        r = requests.get(url, headers=UA, timeout=300)
        if r.status_code == 404:
            log.info("CNH %s: no publicado todavía", edicion)
            continue
        r.raise_for_status()
        wb = openpyxl.load_workbook(io.BytesIO(r.content), read_only=True, data_only=True)
        estructura = {}
        for f in _cnh_hoja(wb, "ESTRUCTURA FUNCIONAL"):
            estructura[str(f.get("CODCNH"))] = f
        n = 0
        for f in _cnh_hoja(wb, "DIRECTORIO DE HOSPITALES"):
            cod = str(f.get("CODCNH"))
            e = estructura.get(cod, {})
            concierto = e.get("Concierto")
            complejo = f.get("CODIDCOM")
            # 'DC' = el dato está en el complejo hospitalario al que pertenece el hospital.
            if concierto == "DC" and complejo:
                concierto = estructura.get(str(complejo), {}).get("Concierto", concierto)
            n += 1
            yield {"edicion": edicion, "codcnh": cod, "nombre": f.get("Nombre Centro"),
                   "cod_municipio": f.get("Cód. Municipio"), "municipio": f.get("Municipio"),
                   "cod_ccaa": f.get("Cód. CCAA"), "ccaa_cnh": f.get("CCAA"),
                   "camas": None if f.get("CAMAS") in (None, "") else int(float(f.get("CAMAS"))),
                   "cod_clase": f.get("Cód. Clase de Centro"), "clase": f.get("Clase de Centro"),
                   "cod_dependencia": f.get("Cód. Dep. Funcional"),
                   "dependencia": f.get("Dependencia Funcional"),
                   "concierto": concierto, "concierto_original": e.get("Concierto"),
                   "codidcom": complejo, "complejo": f.get("Nombre del Complejo"), "url": url}
        log.info("CNH %s: %d hospitales", edicion, n)


@dlt.source(name="sanidad_privada")
def sanidad_privada():
    @dlt.resource(name="sanidad_privada_egsp", write_disposition="replace")
    def egsp():
        yield from _egsp()

    @dlt.resource(name="ine_sp_cobertura", write_disposition="replace")
    def cobertura():
        yield from _cobertura()

    @dlt.resource(name="ine_sp_epf", write_disposition="replace")
    def epf():
        yield from _epf()

    @dlt.resource(name="sanidad_privada_cnh", write_disposition="replace")
    def cnh():
        yield from _cnh()

    return egsp, cobertura, epf, cnh


if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(message)s")
    from ingestion.destino import pipeline

    print(pipeline("sanidad_privada").run(sanidad_privada()))
