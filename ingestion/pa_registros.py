"""Fuente dlt para cuatro indicadores del índice propio de cumplimiento de
publicidad activa y remisión de información (elaboración propia, ver
docs/elaboracion-propia/). Cada recurso guarda el dato oficial tal cual; los
criterios (plazo, estado, universo, foral) se aplican en los modelos dbt
transform/models/marts/pa_ind_*.sql.

1) boe_leyes_ccaa  -> pa_ind_presupuestos_ccaa
   Leyes autonómicas publicadas en la sección I del BOE (sumarios diarios de la
   API de datos abiertos, https://www.boe.es/datosabiertos/api/boe/sumario/AAAAMMDD,
   la misma descarga que ingestion/transparencia_gobierno.py, cuyas funciones se
   reutilizan). Se guardan TODAS las leyes de los departamentos autonómicos de la
   sección I (unas 200-300 al año), no solo las de presupuestos: el criterio de
   «ley de presupuestos del ejercicio» se aplica en dbt y así se puede revisar.
   Para las que tienen «presupuest» en el título se pide además el documento
   (https://www.boe.es/diario_boe/xml.php?id=BOE-A-...) para leer la fecha de la
   disposición (fecha_disposicion), la entrada en vigor (fecha_vigencia) y la
   nota «Publicada en el "Boletín Oficial de ..." número N, de D de MES de AAAA»
   (fecha de publicación en el diario autonómico).
   Caché: data/boe_cache/ccaa/AAAA.json (años cerrados, solo lo filtrado) y
   data/boe_cache/ccaa_docs/<id>.json. Carga incremental con merge por
   identificador; la primera carga recorre desde el 1-sep-1999 (~9.800 días, unos
   20-30 min con 4 hilos).

2) conprel_presupuestos_avance  -> pa_ind_presupuesto
   CONPREL (SGFAL, Hacienda), presupuestos, AVANCE del año en curso para toda
   España (EP{AAAA}CT.xlsx). Columna «Estado Inf. (*)» con valores C / CP / E / N.
   CP = presupuesto prorrogado (deducido, ver docs/elaboracion-propia/
   presupuesto_prorrogado.md). El avance solo se sirve para el ejercicio en curso
   (los anteriores devuelven HTTP 500) y las definitivas pierden la marca CP:
   cada ejecución guarda una FOTO con su fecha (merge por ejercicio + fecha de
   foto + entidad) para ir construyendo histórico. Copia del fichero en
   data/hacienda_cache/conprel_avance/EP{AAAA}CT_{fecha}.xlsx.

3) ispa_sin_remitir  -> pa_ind_ispa
   Información Salarial de Puestos de la Administración (ISPA, Secretaría de
   Estado de Función Pública, digital.gob.es › espacio-ispa): listado anual
   «ayuntamientos_sin_remitir.xlsx» de los ayuntamientos que no han facilitado
   las retribuciones de sus cargos electos. Las rutas cambian cada año: se leen
   de la página «informac_estad_AAAA» de cada edición. Una fila por
   ayuntamiento listado, con la edición, el año de las retribuciones y la
   fecha de la foto (Last-Modified del fichero o, si falta, el día de la
   descarga): «los listados no tienen carácter definitivo». Merge por edición +
   fecha de foto + provincia + municipio.

4) pmp_ccaa  -> pa_ind_pmp_ccaa
   Periodo medio de pago mensual de cada comunidad autónoma (Central de
   Información de Hacienda, libro de series «pmp-series-rd-1040-2017.xlsx»,
   desde abril de 2018). Mide días en PAGAR, no en publicar. Copia en
   data/hacienda_cache/pmp_ccaa/.
"""

import io
import json
import logging
import re
import time
import unicodedata
import xml.etree.ElementTree as ET
from concurrent.futures import ThreadPoolExecutor
from datetime import date, datetime, timedelta
from email.utils import parsedate_to_datetime
from html import unescape
from pathlib import Path
from urllib.parse import urljoin

import dlt
import requests

from ingestion.transparencia_gobierno import _items, _sumario_xml, SUMARIO_URL

log = logging.getLogger(__name__)

REPO_ROOT = Path(__file__).resolve().parent.parent
BOE_CACHE = REPO_ROOT / "data" / "boe_cache" / "ccaa"
BOE_DOCS = REPO_ROOT / "data" / "boe_cache" / "ccaa_docs"
HACIENDA_CACHE = REPO_ROOT / "data" / "hacienda_cache"
UA = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) SpainFacts/1.0 (+https://spainfacts.org)"}

MESES = {
    "enero": 1, "febrero": 2, "marzo": 3, "abril": 4, "mayo": 5, "junio": 6, "julio": 7,
    "agosto": 8, "septiembre": 9, "setiembre": 9, "octubre": 10, "noviembre": 11, "diciembre": 12,
}


def _sin_tildes(s: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFD", s or "") if unicodedata.category(c) != "Mn")


def _get(url: str, intentos: int = 5, **kw) -> requests.Response | None:
    for intento in range(intentos):
        try:
            r = requests.get(url, headers=UA, timeout=120, **kw)
        except requests.RequestException as e:
            log.warning("Error de red en %s (%s)", url, e)
            time.sleep(3 * 2**intento)
            continue
        if r.status_code == 404:
            return None
        if r.status_code in (429,) or r.status_code >= 502:
            time.sleep(3 * 2**intento)
            continue
        return r
    raise RuntimeError(f"No se pudo descargar {url}")


# ---------------------------------------------------------------------------
# 1) Leyes autonómicas en los sumarios del BOE
# ---------------------------------------------------------------------------

BOE_INICIO = date(1999, 9, 1)
BOE_RELECTURA_DIAS = 20
BOE_HILOS = 4

# Departamento del sumario -> código INE de la comunidad (sobre el texto en
# mayúsculas y sin tildes). El orden importa: «CASTILLA-LA MANCHA» antes que
# «CASTILLA Y LEON».
DEP_CCAA = [
    ("ANDALUCIA", "01"), ("ARAGON", "02"), ("ASTURIAS", "03"), ("BALEARS", "04"), ("BALEARES", "04"),
    ("CANARIAS", "05"), ("CANTABRIA", "06"), ("CASTILLA-LA MANCHA", "08"), ("CASTILLA - LA MANCHA", "08"),
    ("CASTILLA LA MANCHA", "08"), ("CASTILLA Y LEON", "07"), ("CATALU", "09"), ("VALENCIA", "10"),
    ("EXTREMADURA", "11"), ("GALICIA", "12"), ("MADRID", "13"), ("MURCIA", "14"), ("NAVARRA", "15"),
    ("PAIS VASCO", "16"), ("EUSKADI", "16"), ("RIOJA", "17"), ("CEUTA", "18"), ("MELILLA", "19"),
]
RE_DEP_AUTONOMICO = re.compile(r"^(COMUNIDAD|COMUNITAT|CIUDAD|REGION|PRINCIPADO|ILLES|ISLAS|JUNTA|GENERALI)")
RE_LEY = re.compile(r"^Ley(es)?\b", re.I)
RE_PRESUP = re.compile(r"presupuest", re.I)
# «Publicada en el BOCYL núm. 249, suplemento, de 29 de diciembre de 2006»,
# «(Publicada íntegramente en el «Boletín Oficial de Cantabria» extraordinario número 50,
# de 29 de diciembre de 2017)», «Publicada en el DOCV núm. 6429, de 31 de diciembre de 2010»
RE_PUBLICADA = re.compile(
    r"Publicad[ao](?:\s+[ií]ntegramente)?\s+en\s+(?:el\s+)?[«\"“]?(?P<diario>[^»\"”,()]{2,120})[»\"”]?"
    r"(?P<extra>[^()]{0,80}?)\bde\s+(?P<d>\d{1,2})\s+de\s+(?P<m>[a-zA-Z]+)\s+de\s+(?P<a>\d{4})",
    re.I,
)
RE_DIARIO_COLA = re.compile(r"\s+(?:n[úu]m(?:ero|\.)?|extraordinario|suplemento)\b.*$", re.I)
RE_NUMERO = re.compile(r"n[úu]m(?:ero|\.)?\s*(\d[\d.]*\d|\d)", re.I)


def _cod_ccaa_dep(dep: str) -> str | None:
    d = _sin_tildes(dep or "").upper()
    if not RE_DEP_AUTONOMICO.search(d):
        return None
    for clave, cod in DEP_CCAA:
        if clave in d:
            return cod
    return None


def _leer_dia_ccaa(dia: date):
    url = SUMARIO_URL.format(dia.strftime("%Y%m%d"))
    for intento in range(6):
        try:
            sumario = _sumario_xml(url)
            break
        except (requests.RequestException, ET.ParseError) as e:
            if intento == 5:
                raise RuntimeError(f"BOE sumario {dia}: {e}") from e
            time.sleep(2 + 4 * intento)
    if sumario is None:
        return []
    filas = []
    for seccion, dep, ep, it in _items(sumario):
        titulo = (it.get("titulo") or "").strip()
        if seccion != "1" or not RE_LEY.search(titulo):
            continue
        cod = _cod_ccaa_dep(dep)
        if not cod:
            continue
        filas.append({
            "identificador": it.get("identificador"),
            "fecha_publicacion_boe": dia.isoformat(),
            "departamento": dep,
            "epigrafe": ep,
            "cod_ccaa": cod,
            "titulo": titulo,
            "url_html": it.get("url_html"),
        })
    return filas


def _leer_anio_boe(anio: int, inicio: date, fin: date, ex: ThreadPoolExecutor) -> list[dict]:
    desde, hasta = max(inicio, date(anio, 1, 1)), min(fin, date(anio, 12, 31))
    completo = desde == date(anio, 1, 1) and hasta == date(anio, 12, 31) and anio < date.today().year
    cache = BOE_CACHE / f"{anio}.json"
    if completo and cache.exists():
        return json.loads(cache.read_text(encoding="utf-8"))
    dias = [desde + timedelta(days=i) for i in range((hasta - desde).days + 1)]
    filas = [f for fs in ex.map(_leer_dia_ccaa, dias) for f in fs]
    if completo:
        BOE_CACHE.mkdir(parents=True, exist_ok=True)
        cache.write_text(json.dumps(filas, ensure_ascii=False), encoding="utf-8")
    return filas


def _fecha8(s: str | None) -> str | None:
    s = (s or "").strip()
    return f"{s[:4]}-{s[4:6]}-{s[6:8]}" if re.fullmatch(r"\d{8}", s) else None


def _doc_boe(identificador: str) -> dict:
    """Metadatos de un documento del BOE (fecha de la disposición, vigencia,
    número oficial y publicación en el diario autonómico)."""
    cache = BOE_DOCS / f"{identificador}.json"
    if cache.exists():
        d = json.loads(cache.read_text(encoding="utf-8"))
        if "publicaciones" in d:
            return _con_publicacion(d)
    r = _get(f"https://www.boe.es/diario_boe/xml.php?id={identificador}")
    if r is None or r.status_code != 200:
        return {}
    raiz = ET.fromstring(r.content)
    meta = raiz.find("metadatos")
    g = (lambda k: (meta.findtext(k) or "").strip()) if meta is not None else (lambda k: "")
    texto = " ".join("".join(raiz.find("texto").itertext()).split()) if raiz.find("texto") is not None else ""
    publicaciones = []
    for m in RE_PUBLICADA.finditer(texto):
        mes = MESES.get(_sin_tildes(m.group("m")).lower())
        if not mes:
            continue
        try:
            fecha = date(int(m.group("a")), mes, int(m.group("d")))
        except ValueError:
            continue
        cola = m.group("diario") + " " + m.group("extra")
        num = RE_NUMERO.search(cola)
        publicaciones.append({
            "diario_autonomico": RE_DIARIO_COLA.sub("", m.group("diario")).strip(" ,«»\"“”"),
            "diario_autonomico_numero": num.group(1).replace(".", "") if num else None,
            "fecha_publicacion_diario": fecha.isoformat(),
        })
    d = {
        "rango": g("rango"),
        "numero_oficial": g("numero_oficial"),
        "fecha_disposicion": _fecha8(g("fecha_disposicion")),
        "fecha_vigencia": _fecha8(g("fecha_vigencia")),
        "departamento_doc": g("departamento"),
        "url_eli": g("url_eli"),
        "publicaciones": publicaciones,
    }
    BOE_DOCS.mkdir(parents=True, exist_ok=True)
    cache.write_text(json.dumps(d, ensure_ascii=False), encoding="utf-8")
    time.sleep(0.3)
    return _con_publicacion(d)


def _con_publicacion(d: dict) -> dict:
    """De las notas «Publicada en…» del texto se queda con la primera cuya fecha cae
    entre la fecha de la ley y 120 días después (el texto cita a menudo otras
    publicaciones: la ley que modifica, convenios, el BOP…)."""
    d = dict(d)
    pubs = d.pop("publicaciones", []) or []
    elegida = None
    if d.get("fecha_disposicion"):
        f0 = date.fromisoformat(d["fecha_disposicion"])
        for p in pubs:
            f = date.fromisoformat(p["fecha_publicacion_diario"])
            if f0 <= f <= f0 + timedelta(days=120):
                elegida = p
                break
    d.update(elegida or {"diario_autonomico": None, "diario_autonomico_numero": None, "fecha_publicacion_diario": None})
    return d


@dlt.resource(name="boe_leyes_ccaa", write_disposition="merge", primary_key="identificador")
def boe_leyes_ccaa(desde: str | None = None, hasta: str | None = None):
    estado = dlt.current.resource_state()
    if desde:
        inicio = date.fromisoformat(desde)
    elif estado.get("ultimo_dia"):
        inicio = date.fromisoformat(estado["ultimo_dia"]) - timedelta(days=BOE_RELECTURA_DIAS)
    else:
        inicio = BOE_INICIO
    fin = date.fromisoformat(hasta) if hasta else date.today()
    log.info("BOE leyes autonómicas: %s -> %s", inicio, fin)
    with ThreadPoolExecutor(max_workers=BOE_HILOS) as ex:
        for anio in range(inicio.year, fin.year + 1):
            filas = _leer_anio_boe(anio, inicio, fin, ex)
            log.info("  %d: %d leyes autonómicas", anio, len(filas))
            for f in filas:
                f = dict(f)
                if RE_PRESUP.search(f["titulo"]):
                    f.update(_doc_boe(f["identificador"]))
                yield f
    estado["ultimo_dia"] = fin.isoformat()


# ---------------------------------------------------------------------------
# 2) CONPREL, avance de presupuestos del año en curso (foto)
# ---------------------------------------------------------------------------

CONPREL_URL = (
    "https://serviciostelematicosext.hacienda.gob.es/SGFAL/CONPREL/Consulta/DescargaFichero"
    "?CCAA=&TipoDato=Presupuestos&Ejercicio={anio}&TipoPublicacion=Avance"
)


def _txt(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return re.sub(r"\s+", " ", str(v)).strip()


def _filas_conprel_avance(contenido: bytes):
    import warnings

    import openpyxl

    # el libro trae áreas de impresión que openpyxl no entiende (avisos inofensivos)
    warnings.filterwarnings("ignore", message="Print area cannot be set", category=UserWarning)
    wb = openpyxl.load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    for ws in wb.worksheets:
        if not ws.title.lower().endswith("entidades"):
            continue
        cols = None
        for fila in ws.iter_rows(values_only=True):
            textos = [_sin_tildes(_txt(x)).lower() for x in fila]
            if cols is None:
                if "pr" in textos and any(t.startswith("estado inf") for t in textos):
                    cols = {
                        "pr": textos.index("pr"),
                        "cor": textos.index("cor"),
                        "tip": next(i for i, t in enumerate(textos) if t.startswith("tip")),
                        "nombre": next(i for i, t in enumerate(textos) if t.startswith("nombre")),
                        "pob": next((i for i, t in enumerate(textos) if t.startswith("pobl")), None),
                        "estado": next(i for i, t in enumerate(textos) if t.startswith("estado inf")),
                    }
                continue
            pr, cor = _txt(fila[cols["pr"]]), _txt(fila[cols["cor"]])
            if not pr.isdigit() or not cor.isdigit():
                continue
            pob = fila[cols["pob"]] if cols["pob"] is not None else None
            yield {
                "hoja": ws.title,
                "cod_prov": pr.zfill(2),
                "cod_corp": cor.zfill(3),
                "tipo": _txt(fila[cols["tip"]]).upper(),
                "nombre": _txt(fila[cols["nombre"]]),
                "poblacion": int(pob) if isinstance(pob, (int, float)) else None,
                "estado_inf": _txt(fila[cols["estado"]]).upper() or None,
            }


@dlt.resource(
    name="conprel_presupuestos_avance",
    write_disposition="merge",
    primary_key=["ejercicio", "fecha_foto", "cod_prov", "cod_corp", "tipo"],
)
def conprel_presupuestos_avance(ejercicio: int | None = None):
    anio = ejercicio or date.today().year
    r = _get(CONPREL_URL.format(anio=anio))
    if r is None or r.status_code != 200 or r.content[:4] != b"PK\x03\x04":
        log.warning("CONPREL: no hay avance de presupuestos %s (HTTP %s)", anio, getattr(r, "status_code", None))
        return
    hoy = date.today()
    destino = HACIENDA_CACHE / "conprel_avance" / f"EP{anio}CT_{hoy.isoformat()}.xlsx"
    destino.parent.mkdir(parents=True, exist_ok=True)
    destino.write_bytes(r.content)
    url = CONPREL_URL.format(anio=anio)
    n = 0
    for f in _filas_conprel_avance(r.content):
        n += 1
        yield {"ejercicio": anio, "fecha_foto": hoy, **f, "url_fuente": url}
    log.info("CONPREL avance %s: %d entidades (foto %s)", anio, n, hoy)


# ---------------------------------------------------------------------------
# 3) ISPA, ayuntamientos que no han facilitado las retribuciones
# ---------------------------------------------------------------------------

ISPA_BASE = "https://digital.gob.es"
ISPA_ESTADISTICAS = ISPA_BASE + "/funcion-publica/dgfp/espacio-ispa/estadisticas"
# La página informac_estad_AAAA es la de las retribuciones del año AAAA, que se
# recogen en la edición ISPA AAAA+1 (informac_estad_2024 -> ispa2025/retrib_2024).
# Las rutas de los Excel cambian casi cada año (ispa_anteriores/ispa_2014/
# ayunt_sinremitir_ispa2014.xlsx, Ispa_2018/retribuc_2017/ayunt_sin_remitir.xlsx,
# ispa2025/retrib_2024/ayuntamientos_sin_remitir.xlsx): se leen de la página.
ISPA_PRIMER_ANIO = 2013


def _url_sin_remitir(anio_retrib: int) -> str | None:
    for pag in (f"{ISPA_ESTADISTICAS}/informac_estad_{anio_retrib}",):
        r = _get(pag)
        if r is None or r.status_code != 200:
            continue
        for h in re.findall(r'href="([^"]+\.xlsx?)"', r.text, flags=re.I):
            h = unescape(h)
            if re.search(r"sin_?remitir|sin-remitir", h, flags=re.I):
                return urljoin(pag, h)
    return None


def _filas_ispa(contenido: bytes):
    if contenido[:4] == b"PK\x03\x04":
        import openpyxl

        wb = openpyxl.load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
        hojas = [(ws.title, list(ws.iter_rows(values_only=True))) for ws in wb.worksheets]
    else:
        import xlrd

        wb = xlrd.open_workbook(file_contents=contenido)
        hojas = [(s.name, [s.row_values(i) for i in range(s.nrows)]) for s in wb.sheets()]
    for hoja, filas in hojas:
        cols = None
        for fila in filas:
            textos = [_sin_tildes(_txt(x)).lower() for x in fila]
            if cols is None:
                if any("municipio" in t or "ayuntamiento" in t or "entidad" in t for t in textos) and any(
                    "provincia" in t for t in textos
                ):
                    cols = {
                        "prov": next(i for i, t in enumerate(textos) if "provincia" in t),
                        "mun": next(
                            i for i, t in enumerate(textos) if "municipio" in t or "ayuntamiento" in t or "entidad" in t
                        ),
                        "ccaa": next((i for i, t in enumerate(textos) if "comunidad" in t or "ccaa" in t), None),
                        "pob": next((i for i, t in enumerate(textos) if "pobla" in t or "habit" in t), None),
                    }
                continue
            mun, prov = _txt(fila[cols["mun"]]), _txt(fila[cols["prov"]])
            if not mun or not prov or mun.lower().startswith("total"):
                continue
            marca = None
            m = re.search(r"\s*\((\*{1,2})\)\s*$", mun)
            if m:
                marca = "remitida_sin_firma" if m.group(1) == "*" else "informacion_dudosa"
                mun = mun[: m.start()].strip()
            pob = fila[cols["pob"]] if cols["pob"] is not None and cols["pob"] < len(fila) else None
            try:
                pob = int(float(pob)) if pob not in (None, "") else None
            except (TypeError, ValueError):
                pob = None
            yield {
                "hoja": hoja,
                "ccaa_nombre": _txt(fila[cols["ccaa"]]) if cols["ccaa"] is not None else None,
                "provincia_nombre": prov,
                "municipio_nombre": mun,
                "poblacion_fichero": pob,
                # (*) «Han remitido la información pero falta la firma»;
                # (**) «Incorpora información dudosa o incompleta, solicitada aclaración»
                "marca": marca,
            }


@dlt.resource(
    name="ispa_sin_remitir",
    write_disposition="merge",
    primary_key=["edicion", "fecha_foto", "provincia_nombre", "municipio_nombre"],
)
def ispa_sin_remitir(anios: list[int] | None = None):
    hoy = date.today()
    for anio_retrib in anios or range(ISPA_PRIMER_ANIO, hoy.year):
        edicion = anio_retrib + 1
        url = _url_sin_remitir(anio_retrib)
        if not url:
            log.info("ISPA %s: sin listado de ayuntamientos sin remitir", edicion)
            continue
        r = _get(url)
        if r is None or r.status_code != 200 or r.content[:4] not in (b"PK\x03\x04", b"\xd0\xcf\x11\xe0"):
            log.warning("ISPA %s: %s no es un Excel", edicion, url)
            continue
        lm = r.headers.get("Last-Modified")
        try:
            foto = parsedate_to_datetime(lm).date() if lm else hoy
        except (TypeError, ValueError):
            foto = hoy
        destino = HACIENDA_CACHE.parent / "ispa_cache" / f"{edicion}_{foto.isoformat()}_{url.rsplit('/', 1)[-1]}"
        destino.parent.mkdir(parents=True, exist_ok=True)
        destino.write_bytes(r.content)
        filas = list(_filas_ispa(r.content))
        log.info("ISPA %s: %d ayuntamientos sin remitir (foto %s, %s)", edicion, len(filas), foto, url)
        vistos = set()
        for f in filas:
            k = (f["provincia_nombre"], f["municipio_nombre"])
            if k in vistos:
                continue
            vistos.add(k)
            yield {
                "edicion": edicion,
                "anio_retribuciones": anio_retrib,
                "fecha_foto": foto,
                "fecha_descarga": hoy,
                **f,
                "url_fuente": url,
            }


# ---------------------------------------------------------------------------
# 4) PMP mensual de las comunidades autónomas (series de Hacienda)
# ---------------------------------------------------------------------------

# Ojo: el fichero1 de PMP_NET NO es el de las comunidades sino el informe
# mensual de las entidades locales del modelo de cesión. El PMP mensual por
# comunidad se publica en la Central de Información de Hacienda
# (hacienda.gob.es › CDI › Periodo medio de pago › PMP de las AAPP) en un libro
# de series mensuales con la metodología del RD 1040/2017 (desde abril de 2018):
# Tablas 4.1 a 4.5 = ratio de operaciones pagadas, importe de pagos realizados
# (la cabecera dice «miles €», pero las magnitudes son euros: Andalucía paga
# unos 600 M€ al mes; se guardan tal cual), ratio de operaciones pendientes,
# importe de pagos pendientes y PMP (días) por comunidad autónoma (17; Ceuta y Melilla
# comunican como entidad local). Códigos S.1312-NN-00-000-B-B-000 con la
# numeración de Hacienda (14 Navarra, 15 País Vasco, 16 La Rioja, 17 C.
# Valenciana), que NO es la del INE: se casa por nombre.
PMP_SERIES_URL = "https://www.hacienda.gob.es/cdi/pmp/pmp-series-rd-1040-2017.xlsx"
PMP_PAGINA = "https://www.hacienda.gob.es/es-es/cdi/paginas/periodomediopago/informacionaapps/pmpdelasaapp.aspx"
NOMBRE_CCAA = [
    ("ANDALUC", "01"), ("ARAG", "02"), ("ASTURIAS", "03"), ("BALEAR", "04"), ("CANARIAS", "05"),
    ("CANTABRIA", "06"), ("MANCHA", "08"), ("LEON", "07"), ("CATALU", "09"), ("VALENCIA", "10"),
    ("EXTREMADURA", "11"), ("GALICIA", "12"), ("MADRID", "13"), ("MURCIA", "14"), ("NAVARRA", "15"),
    ("VASCO", "16"), ("EUSKADI", "16"), ("RIOJA", "17"), ("CEUTA", "18"), ("MELILLA", "19"),
]
TABLAS_PMP = {
    "4.1": "ratio_operaciones_pagadas",
    "4.2": "importe_pagos_realizados",
    "4.3": "ratio_operaciones_pendientes",
    "4.4": "importe_pagos_pendientes",
    "4.5": "pmp_dias",
}


def _cod_ccaa_nombre(nombre: str) -> str | None:
    n = _sin_tildes(nombre).upper()
    for clave, cod in NOMBRE_CCAA:
        if clave in n:
            return cod
    return None


def _mes_columna(v) -> tuple[int, int] | None:
    if isinstance(v, datetime):
        return v.year, v.month
    m = re.fullmatch(r"(\d{4})\s*-\s*(\d{1,2})", _txt(v))
    return (int(m.group(1)), int(m.group(2))) if m else None


def _filas_pmp_series(contenido: bytes):
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    datos: dict[tuple, dict] = {}
    for ws in wb.worksheets:
        tabla = ws.title.replace("Tabla", "").strip()
        if tabla not in TABLAS_PMP:
            continue
        campo = TABLAS_PMP[tabla]
        meses = None
        for fila in ws.iter_rows(values_only=True):
            if meses is None:
                if _txt(fila[0]).lower().startswith("c") and "comunidad" in _sin_tildes(_txt(fila[1])).lower():
                    meses = {i: _mes_columna(v) for i, v in enumerate(fila) if i >= 2 and _mes_columna(v)}
                continue
            codigo = _txt(fila[0])
            if not codigo.startswith("S.1312-"):
                continue
            cod = _cod_ccaa_nombre(_txt(fila[1]))
            for i, (a, m) in meses.items():
                v = fila[i] if i < len(fila) else None
                if not isinstance(v, (int, float)):
                    continue
                d = datos.setdefault((cod, a, m), {"cod_ccaa": cod, "codigo_hacienda": codigo,
                                                   "ccaa_nombre": _txt(fila[1]), "anio": a, "mes": m})
                d[campo] = float(v)
    yield from datos.values()


@dlt.resource(name="pmp_ccaa", write_disposition="replace")
def pmp_ccaa():
    r = _get(PMP_SERIES_URL)
    if r is None or r.status_code != 200 or r.content[:4] != b"PK":
        raise RuntimeError(f"PMP: no se pudo descargar {PMP_SERIES_URL}")
    hoy = date.today()
    destino = HACIENDA_CACHE / "pmp_ccaa" / f"pmp-series-rd-1040-2017_{hoy.isoformat()}.xlsx"
    destino.parent.mkdir(parents=True, exist_ok=True)
    destino.write_bytes(r.content)
    for f in _filas_pmp_series(r.content):
        yield {**f, "fecha_descarga": hoy, "url_fuente": PMP_SERIES_URL}


@dlt.source(name="pa_registros")
def pa_registros():
    return [boe_leyes_ccaa, conprel_presupuestos_avance, ispa_sin_remitir, pmp_ccaa]


if __name__ == "__main__":
    import sys

    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(message)s")
    from ingestion.destino import pipeline

    # uso: python -m ingestion.pa_registros [recurso ...] [desde=AAAA-MM-DD] [hasta=AAAA-MM-DD]
    # (desde/hasta solo para boe_leyes_ccaa: relee los sumarios de ese tramo)
    args = dict(a.split("=", 1) for a in sys.argv[1:] if "=" in a)
    recursos = [a for a in sys.argv[1:] if "=" not in a] or None
    fuente = pa_registros()
    if args:
        fuente.resources["boe_leyes_ccaa"].bind(**args)
    if recursos:
        fuente = fuente.with_resources(*recursos)
    print(datetime.now(), pipeline("pa_registros").run(fuente))
