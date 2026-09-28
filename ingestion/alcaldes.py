"""Fuente dlt para los alcaldes de todos los municipios (1979-hoy).

El Ministerio de Política Territorial y Memoria Democrática publica, en su
Sistema de Información Local (SIL), un Excel por mandato con los alcaldes
de cada ayuntamiento, incluidos los cambios a mitad de mandato (mociones de
censura, dimisiones, fallecimientos). El mandato en curso no está en esos
ficheros: se genera al vuelo en el portal de concejales (redsara.es) y es
una foto del alcalde actual de cada municipio.

No hay API: se descargan los .xlsx y se leen con openpyxl. Los ficheros no
son homogéneos (fila de cabecera distinta, códigos INE como texto o como
número que pierde el cero inicial, fechas como texto dd/mm/aaaa o como fecha
de Excel, nombres "APELLIDOS, NOMBRE" o "NOMBRE APELLIDOS"), así que la
cabecera se localiza buscando la columna del código INE y el resto de campos
se normalizan aquí. La etiqueta del partido se guarda tal cual
(`partido_original`); la agrupación en familias se hace en dbt con el seed
partidos_familias.

Carga completa ("replace"): ~110k filas en total, idempotente y sin estado.

Fuentes:
https://mptmd.gob.es/politica-territorial/local/sistema_de_informacion_local_-SIL-/alcaldes_y_concejales.html
https://concejales.redsara.es/consulta/getAlcaldesLegislatura
"""

import io
import re
import unicodedata
from datetime import date, datetime, timedelta

import dlt
import requests
from openpyxl import load_workbook

ACTUALES_URL = "https://concejales.redsara.es/consulta/getAlcaldesLegislatura"
MANDATO_ACTUAL = "2023-2027"

HISTORICO_URL = (
    "https://mptmd.gob.es/content/dam/mpt/politica-territorial/local/"
    "sistema_de_informacion_local_-SIL-/alcaldes_y_concejales/Alcaldes_Mandato_{inicio}_{fin}.xlsx"
)
# Mandatos cerrados publicados por el SIL (elecciones municipales cada 4 años).
MANDATOS = [(a, a + 4) for a in range(1979, 2023, 4)]

CABECERAS = {"User-Agent": "SpainFacts/1.0 (+https://spainfacts.github.io)"}


def _clave(texto) -> str:
    """Texto en mayúsculas, sin tildes ni signos: para comparar cabeceras."""
    t = unicodedata.normalize("NFKD", str(texto or ""))
    t = "".join(c for c in t if not unicodedata.combining(c))
    return re.sub(r"[^A-Z0-9 ]", "", t.upper()).strip()


def _texto(valor) -> str | None:
    if valor is None:
        return None
    t = re.sub(r"\s+", " ", str(valor)).strip()
    return t or None


def _codigo_ine(valor) -> tuple[str | None, str | None]:
    """Devuelve (código INE de 6 dígitos o None, código de municipio de 5).

    El SIL mezcla tres formatos según el mandato:
      - texto de 6 dígitos "040010" (5 del municipio + dígito de control);
      - número 40010 / 462613: el mismo código de 6 que ha perdido el cero
        inicial de las provincias 01-09 (mandato 1987-1991);
      - texto de 5 dígitos "04001" (1991-2007): código de municipio sin
        dígito de control.
    """
    if valor is None or valor == "":
        return None, None
    if isinstance(valor, (int, float)):
        t = str(int(valor)).zfill(6)
    else:
        t = re.sub(r"\D", "", str(valor))
        if len(t) == 5:
            return None, t
        t = t.zfill(6)
    return (t, t[:5]) if len(t) == 6 else (None, None)


def _fecha(valor) -> date | None:
    if valor is None or valor == "":
        return None
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    if isinstance(valor, (int, float)):
        # número de serie de Excel (días desde 1899-12-30)
        return date(1899, 12, 30) + timedelta(days=int(valor))
    t = str(valor).strip()
    for formato in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%d/%m/%y", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(t, formato).date()
        except ValueError:
            pass
    return None


def _nombre_completo(texto: str | None) -> str | None:
    # "APELLIDO1 APELLIDO2, NOMBRE" -> "NOMBRE APELLIDO1 APELLIDO2"
    if not texto:
        return None
    if "," in texto:
        apellidos, _, nombre = texto.partition(",")
        texto = f"{nombre.strip()} {apellidos.strip()}"
    return re.sub(r"\s+", " ", texto).strip() or None


def _descargar(url: str) -> bytes:
    respuesta = requests.get(url, timeout=300, headers=CABECERAS)
    respuesta.raise_for_status()
    return respuesta.content


def _filas(contenido: bytes):
    """Recorre la primera hoja devolviendo dicts {cabecera_normalizada: valor}.

    La cabecera es la primera fila que contiene una columna "CODIGO INE".
    """
    libro = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    hoja = libro[libro.sheetnames[0]]
    columnas = None
    for fila in hoja.iter_rows(values_only=True):
        if columnas is None:
            claves = [_clave(v) for v in fila]
            if "CODIGO INE" in claves:
                columnas = {i: c for i, c in enumerate(claves) if c}
            continue
        registro = {c: fila[i] for i, c in columnas.items() if i < len(fila)}
        if any(v not in (None, "") for v in registro.values()):
            yield registro
    libro.close()


def _registro(r: dict, mandato: str, nombre: str | None) -> dict | None:
    codigo, cod_mun = _codigo_ine(r.get("CODIGO INE"))
    if not cod_mun:
        return None
    return {
        "mandato": mandato,
        "codigo_ine": codigo,
        "cod_mun": cod_mun,
        "municipio": _texto(r.get("MUNICIPIO")),
        "provincia": _texto(r.get("PROVINCIA")),
        "comunidad": _texto(r.get("COMUNIDAD AUTONOMA")),
        "nombre": nombre,
        "partido_original": _texto(r.get("PARTIDO") or r.get("LISTA")),
        "fecha_posesion": _fecha(r.get("FECHA DE POSESION") or r.get("FECHA POSESION")),
    }


@dlt.resource(name="alcaldes_actuales", write_disposition="replace")
def alcaldes_actuales():
    """Alcalde vigente de cada municipio en el mandato en curso (2023-2027)."""
    for r in _filas(_descargar(ACTUALES_URL)):
        partes = [_texto(r.get(c)) for c in ("NOMBRE", "1ER APELLIDO", "2O APELLIDO")]
        registro = _registro(r, MANDATO_ACTUAL, " ".join(p for p in partes if p) or None)
        if registro:
            registro["cargo"] = _texto(r.get("CARGO"))
            yield registro


@dlt.resource(name="alcaldes_historico", write_disposition="replace")
def alcaldes_historico():
    """Todos los alcaldes de los mandatos cerrados 1979-2023, con los cambios a mitad de mandato."""
    for inicio, fin in MANDATOS:
        contenido = _descargar(HISTORICO_URL.format(inicio=inicio, fin=fin))
        for orden, r in enumerate(_filas(contenido)):
            registro = _registro(r, f"{inicio}-{fin}", _nombre_completo(_texto(r.get("NOMBRE"))))
            if registro:
                registro["fecha_baja"] = _fecha(r.get("FECHA BAJA"))
                registro["orden_fichero"] = orden
                yield registro


@dlt.source(name="alcaldes")
def alcaldes():
    return [alcaldes_actuales, alcaldes_historico]
