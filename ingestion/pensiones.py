"""Fuente dlt de pensiones contributivas, afiliación a la Seguridad Social y gasto en pensiones.

1) Pensiones contributivas en vigor por comunidad autónoma y provincia (Seguridad Social,
   INSS). Cada mes se publica un XLSX «CA<aaaamm>.xlsx» con el número de pensiones, el
   importe y la pensión media por clase (incapacidad permanente, jubilación, viudedad,
   orfandad, favor de familiares). El del último mes está en
     https://www.seg-social.es/wps/portal/wss/internet/EstadisticasPresupuestosEstudios/Estadisticas/EST23/EST24
   y los anteriores (desde enero de 2008) en el histórico
     .../Estadisticas/EST23/2575  (una página por año con subpáginas por tema).
   Los enlaces a los ficheros llevan identificadores que cambian, así que se rastrean
   las páginas. El formato del libro ha cambiado varias veces (nombres de CCAA en
   mayúsculas, columnas Número/P. media o Número/Importe/P. media, orden de clases...):
   se localizan las columnas por sus cabeceras. Solo se lee la hoja «Total sistema».

2) Afiliados medios mensuales a la Seguridad Social (Seguridad Social):
   - serie por regímenes del total del sistema desde enero de 2001
     («19_Serie afiliación media por regímenes (Total Sistema).xlsx»);
   - por provincia y régimen desde enero de 2021 (libros anuales «Afiliación <año>.xlsx»,
     hoja Tabla_3_6 / DATOS PEST1 con las columnas PERIODO, PROVINCIA, DesReg_RETA_SETA,
     Descripcion_Sexo, AUXILIAR, SALDOS). Se agrega por mes, provincia y régimen.
   Página: .../Estadisticas/EST8/EST10/EST290/EST291

3) Gasto público en pensiones en % del PIB (Eurostat):
   - gov_10a_exp: gasto de las AAPP (S13) por función COFOG, GF1002 vejez y
     GF1003 supervivientes, % del PIB, todos los países;
   - spr_exp_pens: gasto en pensiones SEEPROS (todas las prestaciones), % del PIB.

Recursos (replace):
  ss_pensiones_territorio   mes x territorio (España, CCAA, provincia) x clase
  ss_afiliados_regimen      mes x régimen, total del sistema (2001-)
  ss_afiliados_provincia    mes x provincia x régimen (2021-)
  eurostat_pensiones_cofog  año x país x función COFOG, % PIB
  eurostat_pensiones_gasto  año x país, % PIB
"""

import html
import io
import logging
import re
import unicodedata

import dlt
import openpyxl
import requests

from ingestion.eurostat import parse_json_stat_series

log = logging.getLogger(__name__)

SS = "https://www.seg-social.es"
EST = "/wps/portal/wss/internet/EstadisticasPresupuestosEstudios/Estadisticas"
HISTORICO_PENSIONES = EST + "/EST23/2575"
ULTIMO_PENSIONES = EST + "/EST23/EST24"
AFILIACION = EST + "/EST8/EST10/EST290/EST291"
EUROSTAT = "https://ec.europa.eu/eurostat/api/dissemination/statistics/1.0/data/"
CABECERAS = {"User-Agent": "Mozilla/5.0 (spainfacts.org; datos abiertos)"}

_sesion = requests.Session()


def _get(url: str) -> requests.Response:
    r = _sesion.get(url if url.startswith("http") else SS + url, headers=CABECERAS, timeout=180)
    r.raise_for_status()
    return r


def _norm(texto) -> str:
    t = unicodedata.normalize("NFKD", str(texto or "")).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z]", "", t)


# ---------------------------------------------------------------------------
# Territorios: nombre normalizado -> [(nivel, código INE)]
# Las comunidades uniprovinciales se guardan también como provincia.
_TERRITORIOS = {
    "totalsistema": [("pais", "00")], "total": [("pais", "00")], "totalnacional": [("pais", "00")],
    "andalucia": [("ccaa", "01")], "aragon": [("ccaa", "02")],
    "asturias": [("ccaa", "03"), ("provincia", "33")], "asturiasprincipadode": [("ccaa", "03"), ("provincia", "33")],
    "principadodeasturias": [("ccaa", "03"), ("provincia", "33")],
    "ibalears": [("ccaa", "04"), ("provincia", "07")], "balearsilles": [("ccaa", "04"), ("provincia", "07")],
    "illesbalears": [("ccaa", "04"), ("provincia", "07")], "baleares": [("ccaa", "04"), ("provincia", "07")],
    "balears": [("ccaa", "04"), ("provincia", "07")], "ibaleares": [("ccaa", "04"), ("provincia", "07")],
    "canarias": [("ccaa", "05")], "cantabria": [("ccaa", "06"), ("provincia", "39")],
    "castillayleon": [("ccaa", "07")], "castillalamancha": [("ccaa", "08")], "cataluna": [("ccaa", "09")],
    "cvalenciana": [("ccaa", "10")], "comunitatvalenciana": [("ccaa", "10")], "comunidadvalenciana": [("ccaa", "10")],
    "extremadura": [("ccaa", "11")], "galicia": [("ccaa", "12")],
    "madrid": [("ccaa", "13"), ("provincia", "28")], "madridcomde": [("ccaa", "13"), ("provincia", "28")],
    "comunidaddemadrid": [("ccaa", "13"), ("provincia", "28")],
    "murcia": [("ccaa", "14"), ("provincia", "30")], "murciaregionde": [("ccaa", "14"), ("provincia", "30")],
    "regiondemurcia": [("ccaa", "14"), ("provincia", "30")],
    "navarra": [("ccaa", "15"), ("provincia", "31")], "navarracomforalde": [("ccaa", "15"), ("provincia", "31")],
    "comunidadforaldenavarra": [("ccaa", "15"), ("provincia", "31")],
    "paisvasco": [("ccaa", "16")],
    "riojala": [("ccaa", "17"), ("provincia", "26")], "larioja": [("ccaa", "17"), ("provincia", "26")],
    "rioja": [("ccaa", "17"), ("provincia", "26")],
    "ceuta": [("ccaa", "18"), ("provincia", "51")], "melilla": [("ccaa", "19"), ("provincia", "52")],
    # provincias
    "almeria": [("provincia", "04")], "cadiz": [("provincia", "11")], "cordoba": [("provincia", "14")],
    "granada": [("provincia", "18")], "huelva": [("provincia", "21")], "jaen": [("provincia", "23")],
    "malaga": [("provincia", "29")], "sevilla": [("provincia", "41")],
    "huesca": [("provincia", "22")], "teruel": [("provincia", "44")], "zaragoza": [("provincia", "50")],
    "palmaslas": [("provincia", "35")], "laspalmas": [("provincia", "35")],
    "sctenerife": [("provincia", "38")], "santacruzdetenerife": [("provincia", "38")], "stacruztenerife": [("provincia", "38")],
    "sctdetenerife": [("provincia", "38")], "stacruzdetenerife": [("provincia", "38")],
    "avila": [("provincia", "05")], "burgos": [("provincia", "09")], "leon": [("provincia", "24")],
    "palencia": [("provincia", "34")], "salamanca": [("provincia", "37")], "segovia": [("provincia", "40")],
    "soria": [("provincia", "42")], "valladolid": [("provincia", "47")], "zamora": [("provincia", "49")],
    "albacete": [("provincia", "02")], "ciudadreal": [("provincia", "13")], "cuenca": [("provincia", "16")],
    "guadalajara": [("provincia", "19")], "toledo": [("provincia", "45")],
    "barcelona": [("provincia", "08")], "girona": [("provincia", "17")], "gerona": [("provincia", "17")],
    "lleida": [("provincia", "25")], "lerida": [("provincia", "25")], "tarragona": [("provincia", "43")],
    "alicante": [("provincia", "03")], "alicantealacant": [("provincia", "03")],
    "castellon": [("provincia", "12")], "castelloncastello": [("provincia", "12")],
    "valencia": [("provincia", "46")], "valenciavalencia": [("provincia", "46")],
    "badajoz": [("provincia", "06")], "caceres": [("provincia", "10")],
    "corunaa": [("provincia", "15")], "acoruna": [("provincia", "15")], "lacoruna": [("provincia", "15")],
    "coruna": [("provincia", "15")], "corunala": [("provincia", "15")],
    "lugo": [("provincia", "27")], "ourense": [("provincia", "32")], "orense": [("provincia", "32")],
    "pontevedra": [("provincia", "36")],
    "arabaalava": [("provincia", "01")], "alava": [("provincia", "01")], "araba": [("provincia", "01")],
    "gipuzkoa": [("provincia", "20")], "guipuzcoa": [("provincia", "20")],
    "bizkaia": [("provincia", "48")], "vizcaya": [("provincia", "48")],
}
# filas agregadas que no son un territorio único (se ignoran)
_IGNORAR = {"ceutaymelilla", "extranjero", "noconsta", "residentesenelextranjero"}

_CLASES = [
    ("incapacidad", "Incapacidad permanente"), ("jubilac", "Jubilación"), ("viudedad", "Viudedad"),
    ("orfandad", "Orfandad"), ("favor", "Favor de familiares"), ("total", "Total"),
]

_MESES = {"enero": 1, "febrero": 2, "marzo": 3, "abril": 4, "mayo": 5, "junio": 6, "julio": 7,
          "agosto": 8, "septiembre": 9, "setiembre": 9, "octubre": 10, "noviembre": 11, "diciembre": 12}


def _fecha_libro(ws, nombre_fichero: str):
    """Fecha de referencia ('a 1 de <mes> de <año>') del título; si no, del nombre del fichero."""
    for fila in ws.iter_rows(min_row=1, max_row=4, values_only=True):
        for v in fila:
            m = re.search(r"(\d{1,2})\s+de\s+([a-záéíóú]+)\s+(?:de\s+)?(\d{4})", str(v or "").lower())
            if m and _norm(m.group(2)) in _MESES:
                return int(m.group(3)), _MESES[_norm(m.group(2))]
    d = re.sub(r"\D", "", nombre_fichero)
    if len(d) == 7:
        d = d[1:]
    if d[:2] == "20" and 1 <= int(d[4:6]) <= 12:
        return int(d[:4]), int(d[4:6])
    return int(d[2:6]), int(d[:2])


def _hoja_total(wb):
    # la hoja cuyo encabezado dice «TOTAL SISTEMA» (algunos libros de 2022 traen
    # además una hoja «TOTAL INSS» con fórmulas rotas)
    for ws in wb.worksheets:
        for fila in ws.iter_rows(min_row=1, max_row=5, values_only=True):
            if any(_norm(v) == "totalsistema" for v in fila):
                return ws
    for ws in wb.worksheets:
        t = _norm(ws.title)
        if "indice" in t:
            continue
        if "total" in t or "sistema" in t:
            return ws
    return [ws for ws in wb.worksheets if "indice" not in _norm(ws.title)][0]


def _parse_libro_ca(contenido: bytes, nombre_fichero: str):
    wb = openpyxl.load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    ws = _hoja_total(wb)
    anio, mes = _fecha_libro(ws, nombre_fichero)
    filas = list(ws.iter_rows(values_only=True))
    # fila de clases: la que contiene 'jubilación'; fila de medidas: la que contiene 'número'
    i_clases = next(i for i, f in enumerate(filas[:15]) if any("jubilac" in _norm(v) for v in f))
    i_medidas = next(i for i, f in enumerate(filas[:15]) if i >= i_clases and any(_norm(v) == "numero" for v in f))
    clases_fila = filas[i_clases]
    medidas = filas[i_medidas]
    # clase de cada columna: la última cabecera de clase a su izquierda
    col_clase, actual = {}, None
    for j in range(len(medidas)):
        v = clases_fila[j] if j < len(clases_fila) else None
        if v is not None and _norm(v):
            actual = next((c for k, c in _CLASES if k in _norm(v)), None)
        if actual:
            col_clase[j] = actual
    columnas = {}  # clase -> {'numero': j, 'media': j, 'importe': j}
    for j, v in enumerate(medidas):
        n = _norm(v)
        if j not in col_clase or not n:
            continue
        if n.startswith("numero"):
            columnas.setdefault(col_clase[j], {})["numero"] = j
        elif "media" in n:
            columnas.setdefault(col_clase[j], {})["media"] = j
        elif n.startswith("importe"):
            columnas.setdefault(col_clase[j], {})["importe"] = j
    vistos, sin_mapear = set(), set()
    for f in filas[i_medidas + 1:]:
        if not f or f[0] is None or not isinstance(f[1], (int, float)):
            continue
        clave = _norm(f[0])
        if clave in _IGNORAR:
            continue
        territorios = _TERRITORIOS.get(clave)
        if not territorios:
            sin_mapear.add(str(f[0]))
            continue
        for nivel, cod in territorios:
            if (nivel, cod) in vistos:
                continue
            vistos.add((nivel, cod))
            for clase, cols in columnas.items():
                numero = f[cols["numero"]] if "numero" in cols else None
                media = f[cols["media"]] if "media" in cols else None
                if not isinstance(numero, (int, float)):
                    continue
                media = media if isinstance(media, (int, float)) else None
                yield {"anio": anio, "mes": mes, "nivel": nivel, "cod": cod, "nombre": str(f[0]).strip(),
                       "clase": clase, "pensiones": int(numero), "pension_media": float(media) if media is not None else None,
                       "fichero": nombre_fichero}
    if sin_mapear:
        log.warning("%s: filas sin mapear %s", nombre_fichero, sorted(sin_mapear))
    if not any(nivel == "pais" for nivel, _ in vistos):
        log.warning("%s: sin fila de total nacional", nombre_fichero)


def _enlaces_xlsx(texto: str):
    for m in re.finditer(r'href="(/wps/wcm/connect/wss/[^"]+?/([^/"?]+\.xlsx)(?:\?[^"]*)?)"', texto):
        yield html.unescape(requests.utils.unquote(m.group(2).replace("+", " "))), html.unescape(m.group(1))


def _ficheros_ca():
    """(nombre, url) de todos los libros mensuales de pensiones por CCAA y provincia."""
    raiz = _get(HISTORICO_PENSIONES).text
    anios = re.findall(r'href="(' + re.escape(HISTORICO_PENSIONES) + r'/[^"!/]+)"[^>]*>\s*Año\s+(\d{4})', raiz)
    encontrados = {}
    for url_anio, anio in anios:
        vistos, cola = set(), [(url_anio, 0)]
        while cola:
            u, prof = cola.pop(0)
            if u in vistos:
                continue
            vistos.add(u)
            t = _get(u).text
            for nombre, enlace in _enlaces_xlsx(t):
                if re.match(r"CA2?\W?\d", nombre, re.I):
                    encontrados.setdefault(nombre, enlace)
            if prof < 3:
                for m in re.finditer(r'href="(' + re.escape(url_anio) + r'/[^"!]+)"', t):
                    cola.append((m.group(1), prof + 1))
        log.info("pensiones: %s -> %d libros CA acumulados", anio, len(encontrados))
    for nombre, enlace in _enlaces_xlsx(_get(ULTIMO_PENSIONES).text):
        if re.match(r"CA\d", nombre, re.I):
            encontrados.setdefault(nombre, enlace)
    return encontrados


_REGIMENES = [("hogar", "General: empleados de hogar"), ("agrario", "General: agrario"),
              ("generalr", "General"), ("regimengeneral", "General"), ("autonomos", "Autónomos"),
              ("mar", "Mar"), ("carbon", "Minería del carbón")]


def _regimen(texto: str) -> str:
    n = _norm(texto)
    for k, v in _REGIMENES:
        if k in n:
            return v
    return str(texto)


@dlt.source(name="pensiones")
def pensiones():
    @dlt.resource(name="ss_pensiones_territorio", write_disposition="replace")
    def pensiones_territorio():
        ficheros = _ficheros_ca()
        log.info("pensiones: %d libros CA", len(ficheros))
        meses = set()
        for nombre, enlace in sorted(ficheros.items()):
            try:
                filas = list(_parse_libro_ca(_get(enlace).content, nombre))
            except Exception as e:  # un libro corrupto no debe tumbar la carga
                log.warning("pensiones: no se pudo leer %s: %s", nombre, e)
                continue
            if not filas:
                continue
            clave = (filas[0]["anio"], filas[0]["mes"])
            if clave in meses:  # mismo mes publicado dos veces (histórico y último)
                continue
            meses.add(clave)
            yield filas

    def _pagina_afiliacion():
        return list(_enlaces_xlsx(_get(AFILIACION).text))

    @dlt.resource(name="ss_afiliados_regimen", write_disposition="replace")
    def afiliados_regimen():
        nombre, enlace = next((n, e) for n, e in _pagina_afiliacion() if "serie" in _norm(n))
        wb = openpyxl.load_workbook(io.BytesIO(_get(enlace).content), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        # columnas (fijas desde 2001): 1 RG, 2 SE agrario, 3 SE hogar, 4 RETA, 5 SETA, 6-7 Mar c. ajena/propia,
        # 8 Carbón, 9-10 Hogar continuos/discontinuos (hasta 2012), 11-12 R. E. Agrario c. ajena/propia
        # (hasta 2012/2008), 13 total sistema.
        grupos = {"General": [1, 2, 3, 9, 10, 11], "Autónomos": [4, 5, 12], "Mar": [6, 7],
                  "Minería del carbón": [8], "Total": [13]}
        for f in ws.iter_rows(values_only=True):
            m = re.match(r"([a-záéíóú]+)\s+(\d{4})$", str(f[0] or "").strip().lower())
            if not m or _norm(m.group(1)) not in _MESES or not isinstance(f[13], (int, float)):
                continue
            for regimen, cols in grupos.items():
                valor = sum(float(f[c]) for c in cols if isinstance(f[c], (int, float)))
                yield {"anio": int(m.group(2)), "mes": _MESES[_norm(m.group(1))], "regimen": regimen,
                       "afiliados": valor, "fichero": nombre}

    @dlt.resource(name="ss_afiliados_provincia", write_disposition="replace")
    def afiliados_provincia():
        for nombre, enlace in _pagina_afiliacion():
            if not re.search(r"afiliaci.n\s*20\d\d", nombre, re.I):
                continue
            wb = openpyxl.load_workbook(io.BytesIO(_get(enlace).content), read_only=True, data_only=True)
            hoja = next((wb[h] for h in ("Tabla_3_6", "DATOS PEST1") if h in wb.sheetnames), None)
            if hoja is None:
                log.warning("afiliación: %s sin hoja de provincias", nombre)
                continue
            filas = hoja.iter_rows(values_only=True)
            cab = [str(x or "").upper() for x in next(filas)]
            ip = next(i for i, x in enumerate(cab) if x.startswith("PERIOD"))
            ipr, ir, isal, iaux = cab.index("PROVINCIA"), cab.index("DESREG_RETA_SETA"), cab.index("SALDOS"), cab.index("AUXILIAR")
            suma = {}
            for f in filas:
                if f[ip] is None or not isinstance(f[isal], (int, float)):
                    continue
                periodo = str(f[ip]).strip()[:6]
                clave = (int(periodo[:4]), int(periodo[4:6]), f"{int(f[ipr]):02d}", _regimen(f[ir]),
                         "Cuidadores no profesionales" if "CUIDADOR" in str(f[iaux]).upper() else "Sistema")
                suma[clave] = suma.get(clave, 0.0) + float(f[isal])
            log.info("afiliación: %s -> %d filas", nombre, len(suma))
            for (anio, mes, cod_prov, regimen, tipo), valor in suma.items():
                yield {"anio": anio, "mes": mes, "cod_prov": cod_prov, "regimen": regimen, "tipo": tipo,
                       "afiliados": valor, "fichero": nombre}

    @dlt.resource(name="eurostat_pensiones_cofog", write_disposition="replace")
    def pensiones_cofog():
        r = requests.get(EUROSTAT + "gov_10a_exp", timeout=180, params={
            "unit": "PC_GDP", "sector": "S13", "cofog99": ["GF1002", "GF1003", "GF10"], "na_item": "TE",
            "format": "JSON", "lang": "EN"})
        r.raise_for_status()
        datos = r.json()
        etiquetas = datos["dimension"]["geo"]["category"]["label"]
        for coords, valor in parse_json_stat_series(datos):
            yield {"anio": int(coords["time"]), "geo": coords["geo"], "pais": etiquetas.get(coords["geo"]),
                   "cofog": coords["cofog99"], "pc_pib": float(valor)}

    @dlt.resource(name="eurostat_pensiones_gasto", write_disposition="replace")
    def pensiones_gasto():
        r = requests.get(EUROSTAT + "spr_exp_pens", timeout=180, params={
            "unit": "PC_GDP", "spdepb": "TOTAL", "spdepm": "TOTAL", "format": "JSON", "lang": "EN"})
        r.raise_for_status()
        datos = r.json()
        etiquetas = datos["dimension"]["geo"]["category"]["label"]
        for coords, valor in parse_json_stat_series(datos):
            yield {"anio": int(coords["time"]), "geo": coords["geo"], "pais": etiquetas.get(coords["geo"]),
                   "pc_pib": float(valor)}

    return pensiones_territorio, afiliados_regimen, afiliados_provincia, pensiones_cofog, pensiones_gasto


if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(message)s")
    from ingestion.destino import pipeline

    print(pipeline("pensiones").run(pensiones()))
