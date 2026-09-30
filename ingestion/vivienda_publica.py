"""Fuente dlt para el apartado «Vivienda pública en alquiler» (tema `vivienda_publica`).

1) Ministerio de Vivienda y Agenda Urbana (MIVAU), Observatorio de Vivienda y
   Suelo, «Boletín especial Vivienda Social» (PDF; el último es el de 2024,
   publicado en enero de 2025, con la Encuesta sobre vivienda social de 2023
   enviada a las comunidades y a los ayuntamientos de más de 20.000 habitantes).
   Se localiza en la página del Observatorio y se leen con pdfplumber:
   - Tabla 1.4  Calificaciones provisionales de vivienda protegida por régimen de
                uso (propiedad, alquiler, autopromoción), por comunidad, 2005-2023.
   - Tabla 2.1  Vivienda en alquiler social en la UE (Eurostat, Housing Europe y
                elaboración propia): % de las viviendas principales.
   - Tabla 2.2  Parque de vivienda de titularidad autonómica por régimen de
                tenencia en 2023 (arrendamiento de titularidad pública, en
                colaboración público-privada, con opción de compra, venta, otras).
   - Tabla 2.3  Ídem en 2019 y 2023 (evolución).
   - Tabla 2.8  Parque de los ayuntamientos de más de 20.000 habitantes que
                respondieron (los marcados con * repiten el dato de la encuesta
                de 2019; sin cifras = no respondieron).
   - Estimación nacional (capítulo 2, conclusiones): unas 318.000 viviendas
     públicas en alquiler (197.000 autonómicas y 121.000 municipales escaladas por
     población), el 1,72 % de los hogares del Censo 2021.
   La web del Ministerio rechaza clientes sin User-Agent de navegador.

2) OCDE, Affordable Housing Database, indicador PH4.2 «Social rental housing
   stock» (XLSX): viviendas sociales en alquiler y % del parque total de
   viviendas, hacia 2010 y hacia 2022, más las medias UE y OCDE.
   https://webfs.oecd.org/els-com/Affordable_Housing_Database/PH4-2-Social-rental-housing-stock.xlsx

3) INE, Encuesta de Condiciones de Vida (ECV), tabla 9997: hogares por régimen de
   tenencia de la vivienda y comunidad autónoma (%), 2004-..., con la categoría
   «Alquiler inferior al precio de mercado» (incluye el alquiler social público,
   pero también rentas reducidas pactadas entre particulares o de empresa).

Recursos (replace):
  vp_calificaciones   cod_ccaa, anio, regimen, viviendas
  vp_europa           pais, cod_pais, poblacion, parque_total, viviendas_principales,
                      pct_social, viviendas_sociales, anio_dato
  vp_ccaa_parque      cod_ccaa, anio, arrendamiento_publico, arrendamiento_ppp,
                      arrendamiento, opcion_compra, venta, otras, total
  vp_municipios       cod_prov, cod_ccaa, municipio, poblacion, arrendamiento_publico,
                      arrendamiento_ppp, arrendamiento, opcion_compra, venta, otras,
                      total, origen ('encuesta_2023' | 'boletin_2020' | 'sin_respuesta')
  vp_nacional         concepto, valor, anio
  vp_ocde             pais_fuente, cod_pais, anio, viviendas_sociales, pct_parque, es_media
  ine_ecv_tenencia_ccaa  (formato Tempus: cod_serie, serie, fecha, anyo, valor, secreto)
"""

import io
import logging
import re
import unicodedata

import dlt
import requests

from ingestion.ine import _tabla_resource
from ingestion.vivienda import UNIPROVINCIAL, _cod_provincia

log = logging.getLogger(__name__)

CABECERAS = {"User-Agent": "Mozilla/5.0 (spainfacts.org; datos abiertos)"}
OBSERVATORIO = "https://www.mivau.gob.es/urbanismo-y-suelo/suelo/observatorio-de-vivienda-y-suelo"
BOLETIN_2024 = ("https://www.mivau.gob.es/recursos_mfom/comodin/recursos/"
                "observatoriodeviviendaysueloboletnespecialviviendasocial2024_0.pdf")
OCDE_PH42 = "https://webfs.oecd.org/els-com/Affordable_Housing_Database/PH4-2-Social-rental-housing-stock.xlsx"

# Orden de las comunidades en las tablas del boletín (Tabla 1.4)
ORDEN_CCAA = ["01", "02", "03", "04", "05", "06", "07", "08", "09", "10",
              "11", "12", "13", "14", "15", "16", "17", "18", "19", "00"]

# Países de la Tabla 2.1 (nombres del boletín) -> ISO3
PAISES_ES = {
    "alemania": "DEU", "austria": "AUT", "belgica": "BEL", "bulgaria": "BGR", "chequia": "CZE",
    "chipre": "CYP", "croacia": "HRV", "dinamarca": "DNK", "eslovaquia": "SVK", "eslovenia": "SVN",
    "espana": "ESP", "estonia": "EST", "finlandia": "FIN", "francia": "FRA", "grecia": "GRC",
    "hungria": "HUN", "irlanda": "IRL", "italia": "ITA", "letonia": "LVA", "lituania": "LTU",
    "luxemburgo": "LUX", "malta": "MLT", "paisesbajos": "NLD", "polonia": "POL", "portugal": "PRT",
    "rumania": "ROU", "suecia": "SWE", "ue": "EUU",
}

# Países de la OCDE (nombres del XLSX, sin notas) -> ISO3
PAISES_OCDE = {
    "netherlands": "NLD", "austria": "AUT", "denmark": "DNK", "unitedkingdomengland": "GBR",
    "unitedkingdom": "GBR", "france": "FRA", "ireland": "IRL", "iceland": "ISL", "finland": "FIN",
    "korea": "KOR", "poland": "POL", "slovenia": "SVN", "belgium": "BEL", "norway": "NOR",
    "newzealand": "NZL", "czechia": "CZE", "unitedstates": "USA", "canada": "CAN",
    "australia": "AUS", "japan": "JPN", "hungary": "HUN", "germany": "DEU",
    "slovakrepublic": "SVK", "italy": "ITA", "latvia": "LVA", "israel": "ISR", "spain": "ESP",
    "estonia": "EST", "portugal": "PRT", "lithuania": "LTU", "colombia": "COL",
    "switzerland": "CHE", "malta": "MLT", "luxembourg": "LUX", "eu": "EUU", "oecd": "OED",
    "chile": "CHL", "turkiye": "TUR", "romania": "ROU", "mexico": "MEX", "costarica": "CRI",
    "sweden": "SWE", "greece": "GRC", "bulgaria": "BGR", "croatia": "HRV", "cyprus": "CYP",
}


def _norm(texto) -> str:
    t = unicodedata.normalize("NFKD", str(texto)).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z]", "", t)


def _cod_ccaa(nombre: str):
    """Código INE de comunidad a partir de los nombres (y abreviaturas) del boletín."""
    n = _norm(nombre)
    for clave, cod in [
        ("andaluc", "01"), ("aragon", "02"), ("asturias", "03"), ("balears", "04"), ("baleares", "04"),
        ("canarias", "05"), ("cantabria", "06"), ("castillayleon", "07"), ("castillalamancha", "08"),
        ("cataluna", "09"), ("comunitatvalenciana", "10"), ("comunidadvalenciana", "10"),
        ("comvalenciana", "10"), ("extremadura", "11"), ("galicia", "12"), ("madrid", "13"),
        ("murcia", "14"), ("navarra", "15"), ("paisvasco", "16"), ("rioja", "17"), ("ceuta", "18"),
        ("melilla", "19"), ("total", "00"), ("espana", "00"),
    ]:
        if n.startswith(clave):
            return cod
    return None


def _entero(tok):
    """'1.234' -> 1234; '-', 'n.d.', 'ND' -> None."""
    tok = tok.strip()
    if re.fullmatch(r"\d{1,3}(\.\d{3})+|\d+", tok):
        return int(tok.replace(".", ""))
    return None


def _decimal(tok):
    tok = tok.strip().rstrip("%")
    try:
        return float(tok.replace(".", "").replace(",", "."))
    except ValueError:
        return None


def _get(url: str) -> requests.Response:
    r = requests.get(url, headers=CABECERAS, timeout=300)
    r.raise_for_status()
    return r


def _url_boletin() -> str:
    """Último «Boletín Especial Vivienda Social» enlazado en la página del Observatorio."""
    try:
        html = _get(OBSERVATORIO).text
        enlaces = re.findall(r'href="([^"]+)"[^>]*>\s*Bolet[íi]n Especial Vivienda Social (\d{4})', html)
        if enlaces:
            url, anio = max(enlaces, key=lambda e: int(e[1]))
            log.info("Boletín especial Vivienda Social %s: %s", anio, url)
            return url
    except requests.RequestException as e:
        log.warning("No se pudo leer la página del Observatorio (%s); se usa el boletín 2024", e)
    return BOLETIN_2024


def _paginas_boletin() -> list[str]:
    import pdfplumber

    try:
        contenido = _get(_url_boletin()).content
    except requests.RequestException as e:  # el centro de publicaciones falla a veces (certificado)
        log.warning("Centro de publicaciones inaccesible (%s); se usa la copia de mivau.gob.es", e)
        contenido = b""
    if not contenido.startswith(b"%PDF"):
        # el enlace del Observatorio pasa por el centro de publicaciones; si no
        # devuelve el PDF, se usa la copia de la web del Ministerio
        contenido = _get(BOLETIN_2024).content
    with pdfplumber.open(io.BytesIO(contenido)) as pdf:
        return [p.extract_text() or "" for p in pdf.pages]


def _texto_entre(paginas, inicio: str, fin: str) -> str:
    """Texto desde la página donde aparece `inicio` hasta la primera aparición de `fin`."""
    texto = "\n".join(paginas)
    i = texto.find(inicio)
    if i < 0:
        raise ValueError(f"No se encuentra «{inicio}» en el boletín")
    j = texto.find(fin, i + len(inicio))
    return texto[i:j if j > 0 else None]


# --- Tabla 1.4: calificaciones provisionales por régimen -------------------

_TOK = r"(?:\d{1,3}(?:\.\d{3})+|\d+|-|ND|n\.d\.)"
_FILA_14 = re.compile(rf"(Propiedad|Alquiler|Autoprom\.|Total) ({_TOK}(?: {_TOK}){{19}})\s*$")


def calificaciones(paginas):
    texto = _texto_entre(paginas, "Tabla 1.4.", "Tabla 1.5.")
    bloques = []
    for linea in texto.splitlines():
        m = _FILA_14.search(linea.strip())
        if not m:
            continue
        regimen, valores = m.group(1), m.group(2).split()
        if regimen == "Propiedad":
            bloques.append({})
        bloques[-1][regimen] = valores
    if len(bloques) != len(ORDEN_CCAA):
        raise ValueError(f"Tabla 1.4: {len(bloques)} bloques en lugar de {len(ORDEN_CCAA)}")
    nombres = {"Propiedad": "propiedad", "Alquiler": "alquiler", "Autoprom.": "autopromocion", "Total": "total"}
    filas = []
    for cod, bloque in zip(ORDEN_CCAA, bloques):
        for regimen, valores in bloque.items():
            for k, anio in enumerate(range(2005, 2024)):  # la columna 20 es el total 2005-2023
                filas.append({"cod_ccaa": cod, "anio": anio, "regimen": nombres[regimen],
                              "viviendas": _entero(valores[k])})
    # comprobación: el alquiler de España es la suma de las comunidades
    esp = sum(f["viviendas"] or 0 for f in filas if f["cod_ccaa"] == "00" and f["regimen"] == "alquiler")
    ccaa = sum(f["viviendas"] or 0 for f in filas if f["cod_ccaa"] != "00" and f["regimen"] == "alquiler")
    if abs(esp - ccaa) > 0.02 * esp:
        raise ValueError(f"Tabla 1.4: el alquiler de España ({esp}) no cuadra con la suma ({ccaa})")
    return filas


# --- Tabla 2.1: alquiler social en la UE -----------------------------------

_NUM_ES = r"\d{1,3}(?:\.\d{3})*"
_FILA_21 = re.compile(
    rf"^([A-Za-zÁÉÍÓÚáéíóúñÑ ]+?)(\**)\s+(?:27\s+)?({_NUM_ES}) ({_NUM_ES}) ({_NUM_ES}) (\d+,\d)% ({_NUM_ES}) ")


def europa(paginas):
    texto = _texto_entre(paginas, "Tabla 2.1.", "Fuente:")
    filas = []
    for linea in texto.splitlines():
        m = _FILA_21.match(linea.strip())
        if not m:
            continue
        nombre, notas = m.group(1).strip(), m.group(2)
        cod = PAISES_ES.get(_norm(nombre))
        if not cod:
            raise ValueError(f"Tabla 2.1: país desconocido «{nombre}»")
        filas.append({
            "pais": "UE-27" if cod == "EUU" else nombre, "cod_pais": cod,
            "poblacion": _entero(m.group(3)), "parque_total": _entero(m.group(4)),
            "viviendas_principales": _entero(m.group(5)), "pct_social": _decimal(m.group(6)),
            "viviendas_sociales": _entero(m.group(7)),
            # * = dato de 2017 (sin información de 2023); España: ECV 2023
            "anio_dato": 2017 if notas == "*" else 2023,
        })
    if len(filas) < 25:
        raise ValueError(f"Tabla 2.1: solo {len(filas)} países")
    return filas


# --- Tablas 2.2 y 2.3: parque autonómico -----------------------------------

def parque_ccaa(paginas):
    filas = {}
    # 2.2 (2023): nombre + 13 columnas (uds, % alternos y total)
    texto = _texto_entre(paginas, "Tabla 2.2.", "Fuente:")
    for linea in texto.splitlines():
        m = re.match(r"^(\D+?)\s+((?:[\d.,]+%?\s+){12}[\d.]+)$", linea.strip())
        if not m:
            continue
        cod = _cod_ccaa(m.group(1))
        if not cod or cod == "00":
            continue
        v = [t for t in m.group(2).split() if not t.endswith("%")]
        filas[(cod, 2023)] = {
            "cod_ccaa": cod, "anio": 2023, "arrendamiento_publico": _entero(v[0]),
            "arrendamiento_ppp": _entero(v[1]), "arrendamiento": _entero(v[2]),
            "opcion_compra": _entero(v[3]), "venta": _entero(v[4]), "otras": _entero(v[5]),
            "total": _entero(v[6]),
        }
    # 2.3 (2019 y 2023): nombre + 5 grupos de (2019, 2023, var. uds, var. %)
    texto = _texto_entre(paginas, "Tabla 2.3.", "Fuente:")
    for linea in texto.splitlines():
        limpia = re.sub(r"[<>] 100% \(\*\)", "X", linea.strip())
        m = re.match(r"^(\D+?)\s+((?:\S+\s+){19}\S+)$", limpia)
        if not m:
            continue
        cod = _cod_ccaa(m.group(1))
        if not cod or cod == "00":
            continue
        v = m.group(2).split()
        f19 = {"cod_ccaa": cod, "anio": 2019, "arrendamiento_publico": None, "arrendamiento_ppp": None,
               "arrendamiento": _entero(v[0]), "opcion_compra": _entero(v[4]), "venta": _entero(v[8]),
               "otras": _entero(v[12]), "total": _entero(v[16])}
        filas[(cod, 2019)] = f19
        if (cod, 2023) in filas and filas[(cod, 2023)]["arrendamiento"] != _entero(v[1]):
            log.warning("Tabla 2.3 y 2.2 difieren en el alquiler de %s: %s / %s",
                        cod, _entero(v[1]), filas[(cod, 2023)]["arrendamiento"])
    for anio in (2019, 2023):
        n = sum(1 for k in filas if k[1] == anio)
        if n != 19:
            raise ValueError(f"Parque autonómico {anio}: {n} comunidades en lugar de 19")
    return list(filas.values())


# --- Tabla 2.8: parque municipal -------------------------------------------

_POB = r"\d{1,3}\.\d{3}(?:\.\d{3})?"
_PROVINCIAS_EXTRA = {"palmasdegrancanarialas": "35", "castellodelaplana": "12", "castellondelaplana": "12",
                     "castellodelaplanacastellondela": "12"}


def _es_ruido(linea: str) -> bool:
    return (not linea or linea.startswith(("OBSERVATORIO", "DIRECCIÓN GENERAL", "SECRETARÍA GENERAL", "1. En régimen",
                                           "2. En régimen", "1.1. ", "titularidad", "público-privada", "arrendamiento",
                                           "***", "Población 2023", "Tabla 2.8", "entes dependientes")))


def _columnas_municipio(tokens):
    """13 huecos: (uds, %) x 6 + total. Tolera un % sin su cifra (p. ej. '100% 0% 746')."""
    huecos = []
    i = 0
    while i < len(tokens) and len(huecos) < 13:
        t = tokens[i]
        espera_pct = len(huecos) % 2 == 1 and len(huecos) < 12
        if not espera_pct and t.endswith("%") and len(huecos) < 12:
            huecos.append(None)  # falta la cifra
            continue
        huecos.append(t)
        i += 1
    if len(huecos) != 13 or i != len(tokens):
        return None
    c = [_entero(huecos[k]) if huecos[k] else None for k in (0, 2, 4, 6, 8, 10, 12)]
    return {"arrendamiento_publico": c[0], "arrendamiento_ppp": c[1], "arrendamiento": c[2],
            "opcion_compra": c[3], "venta": c[4], "otras": c[5], "total": c[6]}


def municipios(paginas):
    texto = _texto_entre(paginas, "Tabla 2.8. Parque", "* Los datos mostrados corresponden")
    filas, cod_ccaa, cod_prov, ultimo_texto, pendiente = [], None, None, "", None
    no_leidas = []
    for bruta in texto.splitlines():
        linea = bruta.strip()
        if _es_ruido(linea):
            continue
        if not re.search(r"\d", linea):
            # continuación del nombre de la fila anterior ("...San" / "Vicente del Raspeig")
            if pendiente is not None:
                pendiente["municipio"] = f"{pendiente['municipio']} {linea}".strip()
                pendiente = None
                continue
            ultimo_texto = linea
            if linea.isupper():
                cod_ccaa = _cod_ccaa(linea)
                cod_prov = UNIPROVINCIAL.get(cod_ccaa)
                continue
            prov = _PROVINCIAS_EXTRA.get(_norm(linea)) or _cod_provincia(linea)
            if prov:
                cod_prov = prov
            continue
        pendiente = None
        m = re.match(rf"^(.*?)(\*?)\s*({_POB})(?:\s+(.*))?$", linea)
        if not m or cod_prov is None:
            no_leidas.append(linea)
            continue
        nombre = m.group(1).strip()
        if not nombre:  # el nombre estaba en la línea anterior
            nombre = ultimo_texto
        resto = (m.group(4) or "").split()
        fila = {"cod_prov": cod_prov, "cod_ccaa": cod_ccaa, "municipio": nombre,
                "poblacion": _entero(m.group(3))}
        if not resto:
            fila.update({k: None for k in ("arrendamiento_publico", "arrendamiento_ppp", "arrendamiento",
                                           "opcion_compra", "venta", "otras", "total")})
            fila["origen"] = "sin_respuesta"
        else:
            cols = _columnas_municipio(resto)
            if cols is None:
                no_leidas.append(linea)
                continue
            fila.update(cols)
            fila["origen"] = "boletin_2020" if m.group(2) == "*" else "encuesta_2023"
        filas.append(fila)
        if not m.group(1).strip():
            pendiente = fila
    if no_leidas:
        log.warning("Tabla 2.8: %s líneas sin leer: %s", len(no_leidas), no_leidas[:10])
    if len(filas) < 300:
        raise ValueError(f"Tabla 2.8: solo {len(filas)} municipios")
    return filas


# --- Estimación nacional ----------------------------------------------------

def nacional(paginas):
    texto = re.sub(r"\s+", " ", "\n".join(paginas))
    patrones = {
        "parque_alquiler_publico": r"entorno de las ([\d.]+) viviendas",
        "parque_autonomico": r"unas ([\d.]+) son de titularidad de las comunidades",
        "parque_municipal": r"otras ([\d.]+) viviendas son de titularidad de los ayuntamientos",
        "pct_hogares": r"cobertura a un ([\d,]+) ?%",
        "hogares_censo_millones": r"de los ([\d,]+) millones de hogares",
        "pct_hogares_alquiler_inferior_mercado_ecv": r"existe un ([\d,]+) ?% del parque de viviendas principales",
    }
    filas = []
    for concepto, patron in patrones.items():
        m = re.search(patron, texto)
        if not m:
            raise ValueError(f"No se encuentra la estimación nacional «{concepto}»")
        valor = _decimal(m.group(1)) if "," in m.group(1) else float(_entero(m.group(1)))
        filas.append({"concepto": concepto, "valor": valor, "anio": 2023})
    return filas


# --- OCDE PH4.2 --------------------------------------------------------------

def ocde_filas():
    import openpyxl

    libro = openpyxl.load_workbook(io.BytesIO(_get(OCDE_PH42).content), data_only=True, read_only=True)
    filas = []
    for fila in libro["Table PH4.2.A1"].iter_rows(min_row=6, values_only=True):
        nombre = fila[0]
        if not isinstance(nombre, str) or nombre.startswith(("Notes", "Disclaimer", "Source", "The statistical")):
            continue
        base = re.sub(r"\(\d+\)", "", nombre).strip()
        cod = PAISES_OCDE.get(_norm(base))
        if not cod:
            continue
        for n, pct, anio in ((fila[1], fila[2], fila[3]), (fila[4], fila[5], fila[6])):
            if isinstance(pct, (int, float)) and isinstance(anio, (int, float)):
                filas.append({"pais_fuente": base, "cod_pais": cod, "anio": int(anio),
                              "viviendas_sociales": float(n) if isinstance(n, (int, float)) else None,
                              "pct_parque": float(pct), "es_media": False})
    # medias UE y OCDE (último dato de cada país): solo en la figura PH4.2.1
    for fila in libro["Figure PH4.2.1"].iter_rows(min_row=6, values_only=True):
        for k, v in enumerate(fila[:-1]):
            if isinstance(v, str) and v.strip() in ("EU", "OECD") and isinstance(fila[k + 1], (int, float)):
                filas.append({"pais_fuente": v.strip(), "cod_pais": PAISES_OCDE[_norm(v)],
                              "anio": 2022, "viviendas_sociales": None, "pct_parque": float(fila[k + 1]),
                              "es_media": True})
    if not any(f["cod_pais"] == "ESP" for f in filas):
        raise ValueError("OCDE PH4.2 sin dato de España")
    return filas


@dlt.source(name="vivienda_publica")
def vivienda_publica():
    cache = {}

    def paginas():
        if "p" not in cache:
            cache["p"] = _paginas_boletin()
        return cache["p"]

    @dlt.resource(name="vp_calificaciones", write_disposition="replace")
    def r_calificaciones():
        yield calificaciones(paginas())

    @dlt.resource(name="vp_europa", write_disposition="replace")
    def r_europa():
        yield europa(paginas())

    @dlt.resource(name="vp_ccaa_parque", write_disposition="replace")
    def r_ccaa():
        yield parque_ccaa(paginas())

    @dlt.resource(name="vp_municipios", write_disposition="replace")
    def r_municipios():
        yield municipios(paginas())

    @dlt.resource(name="vp_nacional", write_disposition="replace")
    def r_nacional():
        yield nacional(paginas())

    @dlt.resource(name="vp_ocde", write_disposition="replace")
    def r_ocde():
        yield ocde_filas()

    return [r_calificaciones, r_europa, r_ccaa, r_municipios, r_nacional, r_ocde,
            _tabla_resource("ine_ecv_tenencia_ccaa", "9997")]


if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(message)s")
    from ingestion.destino import pipeline

    print(pipeline("vivienda_publica").run(vivienda_publica()))
