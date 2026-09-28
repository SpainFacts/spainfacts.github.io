"""Fuente dlt para las cuentas de los ayuntamientos (Ministerio de Hacienda, CONPREL).

CONPREL es la base de datos de presupuestos y liquidaciones de las entidades
locales de la Secretaría General de Financiación Autonómica y Local:
https://serviciostelematicosext.hacienda.gob.es/SGFAL/CONPREL

No hay API, pero las descargas son GET simples, sin sesión ni cookie:
  .../SGFAL/CONPREL/Consulta/DescargaFichero?CCAA=..&TipoDato=Liquidaciones&Ejercicio=AAAA&TipoPublicacion=..
Si el fichero no existe el servidor responde HTTP 500 con una página de error.

Se usan dos publicaciones:

(B) Resumen por comunidad y ejercicio, liquidación DEFINITIVA
    (`TipoPublicacion=Definitiva&CCAA=01..19`, EL{AAAA}C{cc}.xls, .xlsx desde
    ~2024) y el AVANCE del último ejercicio para toda España
    (`TipoPublicacion=Avance&CCAA=`, EL{AAAA}CT.xlsx; solo se sirve el año en
    curso). La hoja de entidades ("<Región>-Entidades" en el avance, la última
    hoja con el nombre de la comunidad en la definitiva) trae una fila por
    entidad con 33 columnas fijas (32 hasta 2012, sin el capítulo 5 de gastos,
    Fondo de contingencia, y con tipos de una letra: 'A', 'D'):
      Pr, Cor (códigos INE de provincia y municipio), Tipo (A00 ayuntamiento,
      D00 diputación/cabildo/consell, M00 mancomunidad, Z00 Ceuta/Melilla...),
      Nombre, Población, Estado Inf. (C completa, E solo económica, N sin datos),
      ingresos por capítulo 1-9 + total (derechos reconocidos netos),
      gastos por capítulo 1-9 + total (obligaciones reconocidas netas),
      gastos por área de gasto 0, 1, 2, 3, 4, 9 + total.
    Importes en euros. En los .xls antiguos los códigos vienen como números
    (4.0, 1.0) y se rellenan con ceros. Las filas con estado N vienen a cero:
    se cargan tal cual y es dbt quien las convierte en "sin datos" (nulos).
    Ceuta y Melilla (Z00) solo aparecen en el avance nacional: la definitiva
    de las ciudades autónomas no tiene hoja de entidades.
    La numeración de comunidades del parámetro CCAA es la de Hacienda (10
    Extremadura, 11 Galicia, 12 Madrid... distinta del INE), pero aquí da
    igual: los códigos de cada fila son INE.

    Formato ANCHO (una fila por entidad y ejercicio, ~30 columnas numéricas)
    en vez de largo: ~8.100 ayuntamientos x 16 años son ~130k filas en ancho
    frente a ~3,4 M en largo, y los marts las necesitan en ancho de todos modos.

(A) Base de datos Access completa de un ejercicio (`TipoPublicacion=Access`,
    zip de ~45-55 MB con un .accdb de ~450 MB). Solo se usa para el gasto por
    política de gasto (2 dígitos de la clasificación por programas, Orden
    EHA/3565/2008, comparable desde 2010) de los ayuntamientos, en los
    últimos ejercicios con liquidación definitiva:
      tb_inventario: id -> codbdgel ('28079AA000' = provincia(2) + municipio(3)
                     + tipo(2, AA ayuntamiento, ZZ Ceuta/Melilla) + ordinal(3))
      tb_funcional_cons: id, cdcta (capítulo), cdfgr (área 1 dígito, política
                     2, grupo de programas 3, programa 4), importe (obligaciones
                     reconocidas netas, consolidadas con sus organismos). Cada
                     nivel repite el total, así que solo se suman los códigos
                     de 2 dígitos.
      tb_cuentasProgramas: nombres de áreas y políticas.
    access-parser (Python puro) tarda ~45-60 s en la tabla funcional; textos
    rellenos con espacios e importes como texto. El zip y el .accdb se borran
    al terminar cada año.

Recursos (carga completa "replace"):
  conprel_liquidaciones               ayuntamientos (A00) y Ceuta/Melilla (Z00)
  conprel_liquidaciones_diputaciones  diputaciones, cabildos y consells (D00)
  conprel_politicas                   gasto por política de los ayuntamientos; el
                                      mismo recurso escribe la tabla
                                      conprel_politicas_nombres (áreas y políticas)

Duración aproximada de una carga completa: ~10 min para (B) (~300 ficheros
de 0,1-1,5 MB) y ~2-3 min por año de (A) (descarga + lectura del Access).
"""

import io
import logging
import os
import tempfile
import time
import zipfile
from collections import defaultdict
from datetime import date
from pathlib import Path

import dlt
import requests
import xlrd
from openpyxl import load_workbook

URL = "https://serviciostelematicosext.hacienda.gob.es/SGFAL/CONPREL/Consulta/DescargaFichero"
CABECERAS = {"User-Agent": "SpainFacts/1.0 (+https://spainfacts.github.io)"}
PRIMER_ANIO = 2010  # desde 2010 la clasificación por programas es comparable
COMUNIDADES = [f"{i:02d}" for i in range(1, 20)]  # numeración de Hacienda
PAUSA_S = 0.5  # cortesía entre peticiones (todo es secuencial)
ANIOS_POLITICAS_POR_DEFECTO = 3

TIPOS_AYUNTAMIENTO = {"A00", "Z00"}
TIPOS_DIPUTACION = {"D00"}

# Columnas numéricas de la hoja de entidades, por posición (a partir de la 7ª)
COLUMNAS = (
    [f"ingresos_c{i}" for i in range(1, 10)] + ["ingresos_total"]
    + [f"gastos_c{i}" for i in range(1, 10)] + ["gastos_total"]
    + [f"gasto_area_{a}" for a in (0, 1, 2, 3, 4, 9)] + ["gasto_area_total"]
)
# Comprobación de que las columnas no se han movido: (posición, texto que debe contener).
# Solo avisa: las cabeceras traen erratas ("Nombne" en País Vasco 2013).
CABECERAS_ESPERADAS = {6: "impuesto", 16: "personal"}


def _posiciones(cabecera: list[str]) -> list[int | None]:
    """Posición de cada columna de COLUMNAS en la hoja.

    Desde 2013 hay 33 columnas; hasta 2012 no existe el capítulo 5 de gastos
    (Fondo de contingencia) y todo lo que va detrás se desplaza una posición.
    """
    while cabecera and not cabecera[-1]:
        cabecera = cabecera[:-1]
    if len(cabecera) not in (32, 33):
        raise ValueError(f"Cabecera inesperada ({len(cabecera)} columnas): {cabecera}")
    posiciones = list(range(6, 16)) + list(range(16, 20))  # ingresos 1-9+total, gastos 1-4
    if len(cabecera) == 33:
        posiciones += list(range(20, 33))
    else:
        posiciones += [None] + list(range(20, 32))
    if "deuda" not in cabecera[posiciones[COLUMNAS.index("gasto_area_0")]]:
        log.warning("Cabecera de áreas inesperada: %s", cabecera)
    return posiciones


# Tipos explícitos: las diputaciones no traen población y dlt no crearía la columna
COLUMNAS_TIPADAS = {"poblacion": {"data_type": "bigint"},
                    **{c: {"data_type": "double"} for c in COLUMNAS}}

log = logging.getLogger(__name__)


# --------------------------------------------------------------------------
# Descargas
# --------------------------------------------------------------------------

def _descarga(params: dict, *, reintentos: int = 4, stream_a: Path | None = None):
    """GET con reintentos. Devuelve bytes (o la ruta si `stream_a`), o None si no hay fichero.

    El servidor contesta 500 cuando el fichero no existe: tras un par de
    reintentos cortos se da por inexistente (no se puede distinguir de un
    fallo transitorio de otro modo).
    """
    errores_500 = 0
    for intento in range(reintentos + 1):
        time.sleep(PAUSA_S if intento == 0 else min(60, 3 * 2**intento))
        try:
            with requests.get(URL, params=params, headers=CABECERAS, timeout=(30, 300),
                              stream=stream_a is not None) as r:
                if r.status_code >= 500 or r.status_code == 429:
                    errores_500 += 1
                    if r.status_code == 500 and errores_500 >= 2:
                        return None
                    log.warning("CONPREL %s: HTTP %s (intento %d)", params, r.status_code, intento + 1)
                    continue
                r.raise_for_status()
                if "html" in r.headers.get("content-type", ""):
                    return None
                if stream_a is None:
                    return r.content
                with open(stream_a, "wb") as f:
                    for trozo in r.iter_content(1 << 20):
                        f.write(trozo)
                return stream_a
        except requests.RequestException as e:
            log.warning("CONPREL %s: %s (intento %d)", params, e, intento + 1)
    raise RuntimeError(f"No se pudo descargar CONPREL {params}")


def _params(anio: int, publicacion: str, ccaa: str = "") -> dict:
    return {"CCAA": ccaa, "TipoDato": "Liquidaciones", "Ejercicio": anio, "TipoPublicacion": publicacion}


def _existe(anio: int, publicacion: str, ccaa: str = "") -> bool:
    """Comprueba si hay fichero sin descargarlo entero (petición en streaming que se corta)."""
    for intento in range(3):
        time.sleep(PAUSA_S if intento == 0 else 5 * intento)
        try:
            with requests.get(URL, params=_params(anio, publicacion, ccaa), headers=CABECERAS,
                              timeout=(30, 120), stream=True) as r:
                if r.status_code == 500:
                    continue  # no existe (o fallo); se reintenta una vez más
                return r.ok and "attachment" in r.headers.get("content-disposition", "")
        except requests.RequestException as e:
            log.warning("CONPREL sondeo %s %s: %s", anio, publicacion, e)
    return False


def ultimo_definitivo() -> int:
    """Último ejercicio con liquidación definitiva publicada (se sondea hacia atrás)."""
    for anio in range(date.today().year - 1, date.today().year - 6, -1):
        if _existe(anio, "Definitiva", "01"):
            return anio
    raise RuntimeError("CONPREL: no se encuentra ninguna liquidación definitiva reciente")


def anios_avance(ultimo_def: int) -> list[int]:
    """Ejercicios posteriores al último definitivo con avance nacional (normalmente uno)."""
    return [a for a in range(ultimo_def + 1, date.today().year + 1) if _existe(a, "Avance")]


# --------------------------------------------------------------------------
# (B) Hojas de entidades
# --------------------------------------------------------------------------

def _texto(v) -> str:
    return str(v if v is not None else "").strip()


def _codigo(v, ancho: int) -> str:
    """'28' / 28.0 / '001' / 1.0 -> texto con ceros a la izquierda."""
    t = _texto(v)
    if not t:
        return ""
    try:
        return f"{int(float(t)):0{ancho}d}"
    except ValueError:
        return t


def _numero(v) -> float | None:
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    return float(v.strip().replace(".", "").replace(",", "."))


def _entero(v) -> int | None:
    n = _numero(v)
    return int(n) if n is not None else None


def _hojas(contenido: bytes):
    """(nombre, iterador de filas) de cada hoja del libro, .xls o .xlsx."""
    if contenido[:2] == b"PK":
        libro = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
        for nombre in libro.sheetnames:
            yield nombre, libro[nombre].iter_rows(values_only=True)
    else:
        libro = xlrd.open_workbook(file_contents=contenido)
        for hoja in libro.sheets():
            yield hoja.name, (hoja.row_values(i) for i in range(hoja.nrows))


def _es_cabecera(fila) -> bool:
    return len(fila) > 1 and _texto(fila[0]).lower() == "pr" and _texto(fila[1]).lower() == "cor"


def _filas_entidades(contenido: bytes, anio: int, provisional: bool, origen: str):
    """Filas de todas las hojas de entidades del libro (las demás se descartan en cuanto
    se ve que en sus 15 primeras filas no está la cabecera Pr/Cor)."""
    for nombre_hoja, filas in _hojas(contenido):
        cabecera = None
        for i, fila in enumerate(filas):
            fila = list(fila)
            if cabecera is None:
                if _es_cabecera(fila):
                    cabecera = [_texto(c).lower() for c in fila]
                    for pos, texto in CABECERAS_ESPERADAS.items():
                        if texto not in cabecera[pos]:
                            log.warning("%s / %s: columna %d = %r, se esperaba %r",
                                        origen, nombre_hoja, pos, cabecera[pos], texto)
                    posiciones = _posiciones(cabecera)
                elif i > 15:
                    break
                continue
            tipo = _texto(fila[2]).upper()
            if len(tipo) == 1:  # hasta 2012: 'A', 'D'... en vez de 'A00', 'D00'
                tipo += "00"
            pr, cor = _codigo(fila[0], 2), _codigo(fila[1], 3)
            if not tipo or not pr.isdigit():
                continue
            fila_out = {
                "anio": anio,
                "provisional": provisional,
                "cod_mun": pr + cor,
                "cod_prov": pr,
                "nombre": " ".join(_texto(fila[3]).split()),
                "tipo": tipo,
                "poblacion": _entero(fila[4]),
                "estado_informacion": _texto(fila[5]).upper() or None,
            }
            for col, pos in zip(COLUMNAS, posiciones):
                # hasta 2012 no hay capítulo 5 de gastos: 0 (no existía), no nulo
                fila_out[col] = _numero(fila[pos]) if pos is not None else 0.0
            yield fila_out


class _Liquidaciones:
    """Descarga cada fichero una sola vez y reparte sus filas entre los dos recursos."""

    def __init__(self, anio_inicio: int):
        self.ultimo_def = ultimo_definitivo()
        self.ficheros = [(a, "Definitiva", c) for a in range(anio_inicio, self.ultimo_def + 1) for c in COMUNIDADES]
        self.ficheros += [(a, "Avance", "") for a in anios_avance(self.ultimo_def)]
        self._cache: dict[tuple, dict[str, list]] = {}
        log.info("CONPREL: definitiva %s-%s, avance %s", anio_inicio, self.ultimo_def,
                 [a for a, p, _ in self.ficheros if p == "Avance"])

    def _lee(self, clave: tuple) -> dict[str, list]:
        anio, publicacion, ccaa = clave
        contenido = _descarga(_params(anio, publicacion, ccaa))
        partes = {"ayuntamientos": [], "diputaciones": []}
        if contenido is None:
            log.warning("CONPREL: sin fichero %s", clave)
            return partes
        for fila in _filas_entidades(contenido, anio, publicacion == "Avance", f"{anio}-{publicacion}-{ccaa}"):
            if fila["tipo"] in TIPOS_AYUNTAMIENTO:
                partes["ayuntamientos"].append(fila)
            elif fila["tipo"] in TIPOS_DIPUTACION:
                partes["diputaciones"].append(fila)
        log.info("CONPREL %s: %d ayuntamientos, %d diputaciones", clave,
                 len(partes["ayuntamientos"]), len(partes["diputaciones"]))
        return partes

    def filas(self, parte: str):
        for clave in self.ficheros:
            if clave not in self._cache:
                self._cache[clave] = self._lee(clave)
            partes = self._cache[clave]
            filas = partes.pop(parte)
            if not partes:
                del self._cache[clave]
            if filas:
                yield filas


# --------------------------------------------------------------------------
# (A) Access: gasto por política
# --------------------------------------------------------------------------

def _lee_access(anio: int) -> tuple[list[dict], dict[str, str]]:
    """Gasto por política (2 dígitos) de cada ayuntamiento y nombres de áreas/políticas."""
    from access_parser import AccessParser

    logging.getLogger("access_parser").setLevel(logging.ERROR)
    with tempfile.TemporaryDirectory(prefix="conprel_") as tmp:
        ruta_zip = _descarga(_params(anio, "Access"), stream_a=Path(tmp) / f"Liquidaciones{anio}.zip")
        if ruta_zip is None:
            log.warning("CONPREL: sin Access para %s", anio)
            return [], {}
        with zipfile.ZipFile(ruta_zip) as z:
            nombre = next(n for n in z.namelist() if n.lower().endswith((".accdb", ".mdb")))
            ruta_db = Path(z.extract(nombre, tmp))
        ruta_zip.unlink()
        t0 = time.time()
        db = AccessParser(str(ruta_db))
        inventario = db.parse_table("tb_inventario")
        municipio_de = {}
        for id_, cod in zip(inventario["id"], inventario["codbdgel"]):
            cod = cod.strip()
            if cod[5:7] in ("AA", "ZZ"):
                municipio_de[id_.strip()] = cod[:5]
        cuentas = db.parse_table("tb_cuentasProgramas")
        nombres = {c.strip(): " ".join(n.split()) for c, n in zip(cuentas["cdfgr"], cuentas["nombre"])
                   if len(c.strip()) <= 2}
        funcional = db.parse_table("tb_funcional_cons")
        del db
    suma: dict[tuple, float] = defaultdict(float)
    for id_, cdfgr, importe in zip(funcional["id"], funcional["cdfgr"], funcional["importe"]):
        cdfgr = cdfgr.strip()
        if len(cdfgr) != 2:
            continue
        cod_mun = municipio_de.get(id_.strip())
        if cod_mun is None:
            continue
        suma[(cod_mun, cdfgr)] += float(importe.strip().replace(",", ".") or 0)
    log.info("CONPREL Access %s: %d filas de política en %.0f s", anio, len(suma), time.time() - t0)
    filas = [{"anio": anio, "cod_mun": m, "cod_politica": p, "importe": round(v, 2)}
             for (m, p), v in sorted(suma.items())]
    return filas, nombres


def _anios_politicas(anios: list[int] | None, ultimo_def: int | None = None) -> list[int]:
    if anios:
        return sorted(anios)
    env = os.environ.get("CONPREL_ANIOS_POLITICAS", "").strip()
    if env:
        return sorted(int(a) for a in env.replace(";", ",").split(",") if a.strip())
    ultimo = ultimo_def or ultimo_definitivo()
    return list(range(ultimo - ANIOS_POLITICAS_POR_DEFECTO + 1, ultimo + 1))


# --------------------------------------------------------------------------
# Fuente
# --------------------------------------------------------------------------

@dlt.source(name="conprel")
def conprel(anio_inicio: int = PRIMER_ANIO, anios_politicas: list[int] | None = None):
    """Liquidaciones de los presupuestos de las entidades locales (Hacienda, CONPREL).

    anio_inicio: primer ejercicio del resumen por entidad (por defecto 2010).
    anios_politicas: ejercicios del detalle por política (por defecto la
        variable CONPREL_ANIOS_POLITICAS, "2022,2023,2024", o los 3 últimos
        con liquidación definitiva).
    """
    liquidaciones = _Liquidaciones(anio_inicio)

    @dlt.resource(name="conprel_liquidaciones", write_disposition="replace", columns=COLUMNAS_TIPADAS)
    def ayuntamientos():
        yield from liquidaciones.filas("ayuntamientos")

    @dlt.resource(name="conprel_liquidaciones_diputaciones", write_disposition="replace",
                  columns=COLUMNAS_TIPADAS)
    def diputaciones():
        yield from liquidaciones.filas("diputaciones")

    @dlt.resource(name="conprel_politicas", write_disposition="replace")
    def politicas():
        """Gasto por política; al final, los nombres del último año van a su propia tabla."""
        nombres_ultimo: dict[str, str] = {}
        for anio in _anios_politicas(anios_politicas, liquidaciones.ultimo_def):
            filas, nombres = _lee_access(anio)
            nombres_ultimo.update(nombres)  # el último año manda
            if filas:
                yield filas
        yield dlt.mark.with_table_name([
            {
                "codigo": codigo,
                "nivel": "area" if len(codigo) == 1 else "politica",
                "cod_area": codigo[0],
                # Las áreas vienen en mayúsculas: "SERVICIOS PÚBLICOS BÁSICOS"
                "nombre": nombre.capitalize() if nombre.isupper() else nombre,
            }
            for codigo, nombre in sorted(nombres_ultimo.items())
        ], "conprel_politicas_nombres")

    return [ayuntamientos(), diputaciones(), politicas()]
