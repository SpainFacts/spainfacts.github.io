"""Fuente dlt para centrales eléctricas de España (Global Energy Monitor, GIPT).

El Global Integrated Power Tracker de Global Energy Monitor (licencia CC BY 4.0)
es la única fuente con coordenadas, estado (en operación, en construcción,
en tramitación, anunciada, retirada…), potencia, propietario y fechas de cada
central. Su descarga oficial pide rellenar un formulario, así que el fichero no
se descarga automáticamente: se baja a mano y se importa con

    .venv/Scripts/python -m ingestion.gem data/gem/<fichero descargado>.xlsx

que guarda solo las filas de España en ingestion/datos/gem_gipt_espana.csv
(ese extracto sí va al repositorio; el fichero mundial completo, no: data/gem/
está en .gitignore). La carga dlt lee siempre ese extracto.

Cita obligatoria: "Global Integrated Power Tracker, Global Energy Monitor, <edición>".

Al cargar, cada fila recibe el código INE de provincia y de comunidad autónoma
(cod_prov, cod_ccaa) por punto en polígono contra static/geo/provincias.geojson,
porque las columnas de provincia/región del GIPT mezclan inglés, castellano y
lenguas cooficiales. geo_asignacion dice cómo se asignó: 'dentro' (el punto cae
en la provincia), 'cercana' (a menos de 5 km, típicamente en la costa o la raya)
o 'marina' (eólica marina: provincia costera más cercana, hasta 100 km); None si
no se pudo asignar. `python -m ingestion.gem --geo` muestra el resumen.
"""

import csv
import json
import math
import re
import sys
from collections import Counter
from functools import lru_cache
from pathlib import Path

import dlt

EXTRACTO = Path(__file__).resolve().parent / "datos" / "gem_gipt_espana.csv"
PROVINCIAS_GEOJSON = Path(__file__).resolve().parent.parent / "static" / "geo" / "provincias.geojson"
DISTANCIA_MAX_KM = 5.0
DISTANCIA_MAX_MARINA_KM = 100.0


def _normalizar(nombre: str) -> str:
    # "Plant / Project name" y "plant-/-project-name" -> "plant_project_name"
    return re.sub(r"[^a-z0-9]+", "_", str(nombre).strip().lower()).strip("_")


def _es_espana(fila: dict) -> bool:
    pais = str(fila.get("country_area") or fila.get("country") or "")
    return "spain" in pais.lower()


def _leer_descarga(ruta: Path):
    """Filas (dict con columnas normalizadas) de un CSV o de todas las hojas con coordenadas de un Excel."""
    if ruta.suffix.lower() == ".csv":
        with ruta.open(encoding="utf-8-sig", newline="") as f:
            for fila in csv.DictReader(f):
                yield {_normalizar(k): v for k, v in fila.items()}
        return

    from openpyxl import load_workbook

    libro = load_workbook(ruta, read_only=True, data_only=True)
    for hoja in libro.worksheets:
        filas = hoja.iter_rows(values_only=True)
        cabecera = next(filas, None)
        if not cabecera:
            continue
        columnas = [_normalizar(c) for c in cabecera]
        # Las hojas de metadatos/notas no tienen coordenadas
        if not {"latitude", "lat"} & set(columnas):
            continue
        for valores in filas:
            yield {"hoja": hoja.title, **dict(zip(columnas, valores))}


def extraer_espana(ruta_descarga: str) -> int:
    ruta = Path(ruta_descarga)
    filas = [f for f in _leer_descarga(ruta) if _es_espana(f)]
    if not filas:
        raise SystemExit(f"No hay filas de España en {ruta}: ¿es el fichero del GIPT?")
    columnas = sorted({c for f in filas for c in f})
    EXTRACTO.parent.mkdir(parents=True, exist_ok=True)
    with EXTRACTO.open("w", encoding="utf-8", newline="") as f:
        escritor = csv.DictWriter(f, fieldnames=columnas, lineterminator="\n")
        escritor.writeheader()
        escritor.writerows(filas)
    return len(filas)


# --- Provincia de cada central: punto en polígono (Python puro) ------------


@lru_cache(maxsize=1)
def _provincias():
    """[(cod_prov, cod_ccaa, bbox, anillos)] con anillos = listas de (lon, lat)."""
    geo = json.loads(PROVINCIAS_GEOJSON.read_text(encoding="utf-8"))
    provincias = []
    for f in geo["features"]:
        g = f["geometry"]
        poligonos = [g["coordinates"]] if g["type"] == "Polygon" else g["coordinates"]
        # Cada polígono: [exterior, agujeros...]; se guarda como lista de polígonos
        polis = [[[(float(x), float(y)) for x, y in anillo] for anillo in pol] for pol in poligonos]
        xs = [x for pol in polis for x, _ in pol[0]]
        ys = [y for pol in polis for _, y in pol[0]]
        provincias.append(
            (f["properties"]["cod_prov"], f["properties"]["cod_ccaa"], (min(xs), min(ys), max(xs), max(ys)), polis)
        )
    return provincias


def _dentro_anillo(x: float, y: float, anillo) -> bool:
    dentro = False
    j = len(anillo) - 1
    for i in range(len(anillo)):
        xi, yi = anillo[i]
        xj, yj = anillo[j]
        if (yi > y) != (yj > y) and x < (xj - xi) * (y - yi) / (yj - yi) + xi:
            dentro = not dentro
        j = i
    return dentro


def _dentro(x: float, y: float, polis) -> bool:
    for exterior, *agujeros in polis:
        if _dentro_anillo(x, y, exterior) and not any(_dentro_anillo(x, y, a) for a in agujeros):
            return True
    return False


def _distancia_km(lon: float, lat: float, polis) -> float:
    """Distancia (km, proyección equirectangular local) del punto al borde más cercano."""
    kx = 111.32 * math.cos(math.radians(lat))
    ky = 110.57
    mejor = float("inf")
    for pol in polis:
        for anillo in pol:
            for (x1, y1), (x2, y2) in zip(anillo, anillo[1:]):
                ax, ay = (x1 - lon) * kx, (y1 - lat) * ky
                bx, by = (x2 - lon) * kx, (y2 - lat) * ky
                dx, dy = bx - ax, by - ay
                largo2 = dx * dx + dy * dy
                t = 0.0 if largo2 == 0 else max(0.0, min(1.0, -(ax * dx + ay * dy) / largo2))
                mejor = min(mejor, math.hypot(ax + t * dx, ay + t * dy))
    return mejor


@lru_cache(maxsize=None)
def ubicar(lon: float, lat: float, marina: bool = False):
    """(cod_prov, cod_ccaa, asignacion) del punto; (None, None, None) si no cae en España."""
    for cod_prov, cod_ccaa, (x0, y0, x1, y1), polis in _provincias():
        if x0 <= lon <= x1 and y0 <= lat <= y1 and _dentro(lon, lat, polis):
            return cod_prov, cod_ccaa, "dentro"
    limite = DISTANCIA_MAX_MARINA_KM if marina else DISTANCIA_MAX_KM
    margen = limite / 80  # grados, holgado para descartar provincias lejanas por bbox
    candidatas = [
        (_distancia_km(lon, lat, polis), cod_prov, cod_ccaa)
        for cod_prov, cod_ccaa, (x0, y0, x1, y1), polis in _provincias()
        if x0 - margen <= lon <= x1 + margen and y0 - margen <= lat <= y1 + margen
    ]
    if candidatas:
        distancia, cod_prov, cod_ccaa = min(candidatas)
        if distancia <= limite:
            return cod_prov, cod_ccaa, "marina" if marina else "cercana"
    return None, None, None


def _con_provincia(fila: dict) -> dict:
    try:
        lat, lon = float(fila["latitude"]), float(fila["longitude"])
    except (KeyError, TypeError, ValueError):
        return {**fila, "cod_prov": None, "cod_ccaa": None, "geo_asignacion": None}
    marina = "offshore" in str(fila.get("technology") or "").lower()
    cod_prov, cod_ccaa, asignacion = ubicar(round(lon, 5), round(lat, 5), marina)
    return {**fila, "cod_prov": cod_prov, "cod_ccaa": cod_ccaa, "geo_asignacion": asignacion}


def resumen_geo() -> Counter:
    with EXTRACTO.open(encoding="utf-8", newline="") as f:
        filas = [_con_provincia(fila) for fila in csv.DictReader(f)]
    sin = [f for f in filas if f["cod_prov"] is None]
    for f in sin:
        print(
            f"  sin provincia: {f['plant_project_name']} | {f['unit_phase_name']} | "
            f"{f['latitude']},{f['longitude']} | {f['technology']} | {f['subnational_unit_state_province']}"
        )
    return Counter(f["geo_asignacion"] for f in filas)


@dlt.resource(name="gem_centrales", write_disposition="replace")
def gem_centrales():
    if not EXTRACTO.exists():
        raise FileNotFoundError(
            f"Falta {EXTRACTO}: descarga el GIPT de Global Energy Monitor e impórtalo con "
            "`python -m ingestion.gem <fichero>` (ver docstring de ingestion/gem.py)."
        )
    with EXTRACTO.open(encoding="utf-8", newline="") as f:
        for fila in csv.DictReader(f):
            # Vacíos a None para que dlt infiera tipos y no cree columnas de texto vacías
            fila = {k: (v if v not in ("", "not found") else None) for k, v in fila.items()}
            yield _con_provincia(fila)


@dlt.source(name="gem")
def gem():
    return [gem_centrales]


if __name__ == "__main__":
    if sys.argv[1:] == ["--geo"]:
        print(dict(resumen_geo()))
        raise SystemExit(0)
    if len(sys.argv) != 2:
        raise SystemExit("Uso: python -m ingestion.gem <fichero GIPT descargado (.xlsx o .csv)>")
    print(f"{extraer_espana(sys.argv[1])} filas de España guardadas en {EXTRACTO}")
